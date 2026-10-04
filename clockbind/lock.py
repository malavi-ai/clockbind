"""Data lock for the Bridge study: data dictionary + gated lock with hashes.

    data_dictionary(workbook)  -> DataFrame describing every sheet/column
    run_lock(...)              -> refuses unless the workbook is lock-ready; otherwise writes a
                                  read-only lock package, a manifest with SHA-256 hashes and one
                                  lock hash for external registration.

The tool never edits the workbook and never closes a PENDING item or a contradiction itself.
"""
from __future__ import annotations

import csv
import datetime as _dt
import hashlib
import json
import os
import shutil
import stat
from pathlib import Path

import pandas as pd

# ----------------------------------------------------------------- data dictionary
DESCR = {
    "Source code": "Pseudonymised code of one source record (registered in 01).",
    "Company": "Group company that holds the record.",
    "Source class": "Kind of source (email, invoice, bank record, logistics …).",
    "Date range (from–to)": "Period the source covers.",
    "Duplicate of (code)": "Code of the record this one duplicates, if any.",
    "Unique?": "Yes if not a duplicate.",
    "In archive window?": "Yes if inside the analytic archive window.",
    "Access cleared?": "Yes if access is cleared (controller / KVKK); Held if not.",
    "Analytic source?": "Calculated: unique, in window and cleared.",
    "Searched in 02A?": "Calculated: the source appears in the fixed-term search log.",
    "Exclusion reason": "Reason a record is not analytic (rows are never deleted).",
    "Candidate ID": "Episode candidate code (organisation + R + tW).",
    "Organisation code": "Pseudonymised organisation code.",
    "Decision": "Retained / Merged / Not admitted.",
    "Merged into": "Target candidate when Merged.",
    "Episode ID": "Pseudonymised episode code.",
    "Entry kind": "Event, expectation, deadline (tW) or required event (R).",
    "Clock": "Finance, payment, fulfilment, logistics, operational readiness.",
    "Date precision": "Exact, range or not documented.",
    "Expected duration (days, ex-ante)": "Duration stated before the outcome (whole days).",
    "Stated by": "Who wrote the document: firm, customer or third party.",
    "Author side": "Who wrote the document: firm, customer or third party.",
    "Commitment-document date": "Date of the document that fixes the commitment amount (ECB rate date).",
    "HUF equivalent": "Calculated: amount × ECB rate (HUF per unit).",
    "≥ HUF 200,000?": "Calculated materiality floor for L1b.",
    "Calculated level": "Calculated from the gate criteria (never typed).",
    "Final level": "Calculated level unless a logged override exists.",
    "Override level": "Manual override; needs reason and date.",
    "Verdict": "Binding verdict (Level 3 only, after the freeze).",
    "Claim tier": "Calculated claim ceiling (A, B, C, D, E, X).",
    "Permitted wording": "Calculated wording allowed by the claim tier.",
    "Row check": "Calculated integrity check of the row.",
    "Link check": "Calculated: episode is Retained in 02B and not duplicated.",
    "L3e-S source type": "Sensitivity S1/S2 only: type of record dating the customer-side event; never changes the primary level.",
    "L3b gap reason": "Descriptive, only where L3b = No: why the ex-ante duration record is missing. Never changes a gate or level.",
    "L3c gap reason": "Descriptive, only where L3c = No: why an outcome-independent window record is missing. Never changes a gate or level.",
    "L3e gap reason": "Descriptive, only where L3e = No: why an independent customer-side date is missing. Never changes a gate or level.",
    "Stop point": "Descriptive, episodes below Level 1: main reason the episode stopped. 'Stopped after price quote' is not a financing constraint.",
    "Stop point 2": "Descriptive, optional second reason (same list as Stop point).",
    "Buyer stage": "Descriptive, every analysable episode: existing business or new venture (pre-opening / <1 year).",
    "zero_documented": "Steps: Yes only if expected_days = 0 and a pre-outcome document states immediate/same-day execution; otherwise a 0 is treated as missing.",
    "episode": "Episode code (binding input).",
    "anchor_earliest": "Start of the process for steps without predecessors (earliest bound).",
    "tw_earliest": "Deadline tW (earliest bound).",
    "r_step": "Step whose finish is the required event R.",
    "step": "Process step name.",
    "predecessors": "Steps that must finish first (';'-separated).",
    "finish_earliest": "Documented finish date of the step (earliest bound).",
    "finish_latest": "Documented finish date of the step (latest bound; blank if exact).",
    "expected_days": "Ex-ante expected duration of the step (days).",
    "clock": "Clock of the step: finance, payment, fulfilment, logistics, operational-readiness.",
    "source": "Document code supporting the date (no names or amounts).",
    "note": "Free note (no personal data).",
    "anchor_latest": "Start of the process (latest bound; blank if exact).",
    "tw_latest": "Deadline tW (latest bound; blank if exact).",
}


def _header_row(ws) -> int:
    for r in range(1, 7):
        vals = [c.value for c in ws[r]]
        filled = [v for v in vals if v not in (None, "")]
        if len(filled) >= 3 and all(isinstance(v, str) for v in filled):
            return r
    return 1


def _validations(ws) -> dict:
    out = {}
    from openpyxl.utils import range_boundaries
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list":
            f = (dv.formula1 or "").strip('"')
            desc = f"one of: {f}" if not f.startswith(("'", "=")) and "!" not in f else f"code from {f}"
        elif dv.type in ("date", "whole", "decimal"):
            desc = f"{dv.type} {dv.operator or ''} {dv.formula1 or ''} {dv.formula2 or ''}".strip()
        else:
            desc = dv.type or ""
        for rng in str(dv.sqref).split():
            c1, _, c2, _ = range_boundaries(rng)
            for c in range(c1, (c2 or c1) + 1):
                out.setdefault(c, desc)
    return out


def data_dictionary(workbook: str, gates: str | None = None) -> pd.DataFrame:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    labels = {}
    if gates:
        g = json.loads(Path(gates).read_text(encoding="utf-8"))
        for s in g.get("stages", []):
            for c in s.get("criteria", []):
                labels[c["column"]] = f"Gate {c['id']}" + ("" if c.get("gate", True) is not False else " (flag/analysis step)") + (f": {c.get('label', '')}" if c.get("label") else "")
    wb = load_workbook(workbook)
    rows = []
    for ws in wb.worksheets:
        hr = _header_row(ws)
        dvs = _validations(ws)
        for c in range(1, ws.max_column + 1):
            h = ws.cell(hr, c).value
            if h in (None, ""):
                continue
            probe = ws.cell(hr + 1, c)
            kind = "calculated" if isinstance(probe.value, str) and probe.value.startswith("=") else "input"
            rows.append({"sheet": ws.title, "header row": hr, "column": get_column_letter(c), "field": str(h),
                         "entry": kind, "allowed values / validation": dvs.get(c, ""),
                         "number format": probe.number_format if probe.number_format != "General" else "",
                         "meaning": labels.get(str(h)) or DESCR.get(str(h), "")})
    return pd.DataFrame(rows)


# ----------------------------------------------------------------- helpers
def _sha(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _readonly(path: Path):
    for p in [path] + (list(path.rglob("*")) if path.is_dir() else []):
        if p.is_file():
            os.chmod(p, stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)


def _read_notes(path) -> dict:
    if not path:
        return {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {r["check"].strip(): r.get("justification", "").strip() for r in csv.DictReader(f) if r.get("check")}


# ----------------------------------------------------------------- readiness
def readiness(workbook, gates, work_dir, sheet="05_Level_Assessment", header_row=2, notes=None):
    """Returns (blockers, warnings, report_result, eps)."""
    from .fullreport import build
    from .plugins.screen import apply_gates, check_frozen, load_data_sheet, load_gates
    blockers, warnings = [], []
    g = load_gates(gates)
    try:
        _, frozen = check_frozen(g, gates)
    except SystemExit as e:
        frozen = False
        blockers.append(("Gates", str(e)))
    if not frozen and not any(b[0] == "Gates" for b in blockers):
        blockers.append(("Gates", f"{g.get('protocol_version')} is not frozen; lock needs the frozen gates file."))
    r = build(workbook, gates, str(work_dir), sheet=sheet, header_row=header_row, chat_safe=False, title="Bridge — pre-lock report")
    t = r["tables"]
    if "01_Integrity_detail" in t:
        d = t["01_Integrity_detail"]
        for _, x in d[d["level"] == "FAIL"].iterrows():
            blockers.append(("Workbook integrity", f"{x['check']} — {x['where']}"))
        for _, x in d[d["level"] == "WARN"].iterrows():
            warnings.append(("Workbook integrity", f"{x['check']} — {x['where']}"))
    if "02_Screening_checks" in t and len(t["02_Screening_checks"]):
        s = t["02_Screening_checks"]
        lv = s.columns[0]
        for _, x in s[s[lv] == "ERROR"].iterrows():
            blockers.append(("Screening", " — ".join(str(v) for v in x.values[1:3])))
        for _, x in s[s[lv] == "WARN"].iterrows():
            warnings.append(("Screening", " — ".join(str(v) for v in x.values[1:3])))
    c = t.get("11_Consistency_checks")
    if c is not None:
        for _, x in c[c["level"] == "FAIL"].iterrows():
            blockers.append(("Consistency", f"{x['check']} ({x['episodes']})"))
        for _, x in c[c["level"] == "WARN"].iterrows():
            warnings.append(("Consistency", f"{x['check']} ({x['episodes']})"))
    # PENDING in gate criteria of analysable episodes
    df = load_data_sheet(workbook, sheet, header_row)
    _, eps, _ = apply_gates(df, g)
    eid = "Episode ID" if "Episode ID" in eps.columns else eps.columns[0]
    gcols = [c["column"] for s in g["stages"] for c in s["criteria"] if c.get("gate", True) is not False and c["column"] in eps.columns]
    analysable = eps[~eps["level"].astype(str).str.startswith(("Outside", "Held"))]
    pend = [str(e) for e, row in analysable.set_index(eid)[gcols].iterrows() if (row.astype(str).str.strip() == "PENDING").any()]
    if pend:
        blockers.append(("PENDING", f"Gate criteria still PENDING for analysable episodes: {', '.join(pend)} — code from documents or close under the frozen cutoff rule."))
    # documented warnings
    notes = notes or {}
    undocumented = [w for w in warnings if not any(k and k in w[1] for k in notes)]
    for w in undocumented:
        blockers.append(("Undocumented warning", f"{w[0]}: {w[1]} — fix it, or add a justification row to the warnings note (column 'check' must contain this text)."))
    return blockers, warnings, r, eps


# ----------------------------------------------------------------- lock
def run_lock(workbook, gates, by, out_root, sheet="05_Level_Assessment", header_row=2, notes_path=None,
             amendment="", dry_run=False, extra_files=()):
    from . import __version__
    from .core.provenance import code_sha256, package_versions
    out_root = Path(out_root)
    out_root.mkdir(parents=True, exist_ok=True)
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    work = out_root / f"_prelock_{stamp}"
    notes = _read_notes(notes_path)
    blockers, warnings, rep, eps = readiness(workbook, gates, work, sheet, header_row, notes)

    reg_path = out_root / "LOCK_REGISTER.csv"
    previous = []
    if reg_path.exists():
        with open(reg_path, encoding="utf-8") as f:
            previous = list(csv.DictReader(f))
    if previous and not amendment:
        blockers.append(("Amendment", f"A lock already exists ({previous[-1]['lock_id']}). A new lock is an amendment: pass --amendment with the reason."))

    lines = [f"# Lock readiness — {stamp}", "", f"Workbook: {Path(workbook).name}", f"Gates: {Path(gates).name}", "",
             f"**Blockers: {len(blockers)}**  ·  Warnings: {len(warnings)} (documented: {len(warnings) - sum(1 for b in blockers if b[0] == 'Undocumented warning')})", ""]
    if blockers:
        lines += ["## Blockers (lock refused)", ""] + [f"- **{a}:** {b}" for a, b in blockers] + [""]
    if warnings:
        lines += ["## Warnings", ""] + [f"- {a}: {b}" + (f" — *justified:* {next(v for k, v in notes.items() if k and k in b)}" if any(k and k in b for k in notes) else "") for a, b in warnings] + [""]
    readiness_md = work / "Lock_Readiness.md"
    readiness_md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if blockers or dry_run:
        return {"locked": False, "blockers": blockers, "warnings": warnings, "readiness": readiness_md, "work": work}

    # ---- lock package
    lock_id = f"LOCK_{stamp}"
    pkg = out_root / lock_id
    pkg.mkdir()
    shutil.copy2(workbook, pkg / Path(workbook).name)
    shutil.copy2(gates, pkg / Path(gates).name)
    regf = Path(gates).resolve().parent / "FREEZE_REGISTER.jsonl"
    if regf.exists():
        shutil.copy2(regf, pkg / regf.name)
    dd = data_dictionary(workbook, gates)
    with pd.ExcelWriter(pkg / "Data_Dictionary.xlsx", engine="openpyxl") as xw:
        dd.to_excel(xw, sheet_name="Data_Dictionary", index=False)
    dd.to_csv(pkg / "Data_Dictionary.csv", index=False)
    eid = "Episode ID" if "Episode ID" in eps.columns else eps.columns[0]
    held = eps[eps["level"].astype(str).str.startswith(("Outside", "Held"))][[eid, "level"]]
    held.to_csv(pkg / "Held_and_Excluded_Register.csv", index=False)
    shutil.copytree(work, pkg / "locked_report", ignore=shutil.ignore_patterns("_prelock*"))
    if notes_path:
        shutil.copy2(notes_path, pkg / "Warnings_Justification.csv")
    for f in extra_files:
        if f and Path(f).exists():
            shutil.copy2(f, pkg / Path(f).name)
    files = sorted(p for p in pkg.rglob("*") if p.is_file())
    manifest = {
        "lock_id": lock_id, "locked_at": _dt.datetime.now().isoformat(timespec="seconds"), "locked_by": by,
        "amendment_of": previous[-1]["lock_id"] if previous else None, "amendment_reason": amendment or None,
        "workbook": Path(workbook).name, "gates": Path(gates).name,
        "software": {"clockbind": __version__, "code_sha256": code_sha256(), "packages": package_versions()},
        "counts": {"episodes": int(len(eps)), "held_or_excluded": int(len(held)), "warnings_documented": len(warnings)},
        "files": {str(p.relative_to(pkg)): _sha(p) for p in files},
    }
    mtxt = json.dumps(manifest, indent=2, ensure_ascii=False, sort_keys=True)
    (pkg / "Lock_Manifest.json").write_text(mtxt, encoding="utf-8")
    lock_hash = hashlib.sha256(mtxt.encode("utf-8")).hexdigest()
    (pkg / "LOCK_HASH.txt").write_text(f"{lock_id}\nSHA-256 of Lock_Manifest.json: {lock_hash}\n", encoding="utf-8")
    md = [f"# {lock_id}", "", f"- Locked at: {manifest['locked_at']} by {by}", f"- Lock hash (SHA-256 of Lock_Manifest.json): `{lock_hash}`",
          f"- Workbook: {manifest['workbook']} · Gates: {manifest['gates']}", f"- ClockBind {__version__} · code {manifest['software']['code_sha256'][:16]}",
          f"- Episodes: {manifest['counts']['episodes']} · Held/excluded: {manifest['counts']['held_or_excluded']} · documented warnings: {len(warnings)}"]
    if amendment:
        md += [f"- Amendment of {manifest['amendment_of']}: {amendment}"]
    md += ["", "## Files", ""] + [f"- `{k}` — {v[:16]}…" for k, v in manifest["files"].items()]
    md += ["", "## Next", "", "1. Register the lock hash externally (OSF hash-only note, or email to the supervisor) today.",
           f'2. Change Log: "USER DECISION — {manifest["locked_at"][:10]}: data lock {lock_id} (SHA-256 {lock_hash[:12]}). After this point data change only by logged amendment."',
           "3. All manuscript numbers come from locked_report/."]
    (pkg / "Lock_Manifest.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    _readonly(pkg)
    new = not reg_path.exists()
    with open(reg_path, "a", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["lock_id", "locked_at", "locked_by", "lock_hash", "workbook_sha256", "amendment_reason"])
        w.writerow([lock_id, manifest["locked_at"], by, lock_hash, manifest["files"][Path(workbook).name], amendment])
    shutil.rmtree(work, ignore_errors=True)
    return {"locked": True, "lock_id": lock_id, "lock_hash": lock_hash, "package": pkg, "warnings": warnings, "blockers": []}


def verify(package) -> tuple[bool, list]:
    """Re-hash every file in a lock package against its manifest."""
    pkg = Path(package)
    m = json.loads((pkg / "Lock_Manifest.json").read_text(encoding="utf-8"))
    problems = []
    for rel, h in m["files"].items():
        p = pkg / rel
        if not p.exists():
            problems.append(f"missing: {rel}")
        elif _sha(p) != h:
            problems.append(f"changed: {rel}")
    listed = set(m["files"]) | {"Lock_Manifest.json", "Lock_Manifest.md", "LOCK_HASH.txt"}
    problems += [f"added after lock: {p.relative_to(pkg)}" for p in pkg.rglob("*") if p.is_file() and str(p.relative_to(pkg)) not in listed]
    mh = hashlib.sha256((pkg / "Lock_Manifest.json").read_bytes()).hexdigest()
    rec = (pkg / "LOCK_HASH.txt").read_text(encoding="utf-8") if (pkg / "LOCK_HASH.txt").exists() else ""
    if mh not in rec:
        problems.append("Lock_Manifest.json does not match LOCK_HASH.txt")
    return not problems, problems
