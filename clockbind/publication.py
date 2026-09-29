"""Publication and preregistration packaging utilities.

These utilities create immutable *copies* and hash manifests.  They never claim
that a protocol has been registered on OSF or another registry; registration is
an external action that must be performed and evidenced separately.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import zipfile


def sha256_file(path: str | Path) -> str:
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()


def _safe_name(p: Path) -> str:
    return p.name.replace("/","_").replace("\\","_")


def build_prereg_bundle(project_manifest: str | Path | None, files: list[str | Path], output: str | Path,
                         status: str="DRAFT", frozen_by: str="", note: str="") -> Path:
    status=status.upper().strip()
    if status not in {"DRAFT","FROZEN"}: raise ValueError("status must be DRAFT or FROZEN")
    src=[]
    if project_manifest:
        p=Path(project_manifest).expanduser();
        if not p.exists(): raise FileNotFoundError(p)
        src.append(("project_manifest",p))
    for f in files:
        p=Path(f).expanduser()
        if not p.exists(): raise FileNotFoundError(p)
        if p.suffix.lower() in {".csv",".xlsx",".xls",".sav",".parquet",".dta"}:
            raise ValueError(f"Data file refused in preregistration bundle: {p.name}. Include protocols/syntax/codebooks, not row-level research data.")
        src.append(("document",p))
    if not src: raise ValueError("Add at least one protocol/plan/syntax file")
    out=Path(output).expanduser(); out.parent.mkdir(parents=True,exist_ok=True)
    created=dt.datetime.now(dt.timezone.utc).isoformat()
    entries=[]
    with tempfile.TemporaryDirectory() as td:
        root=Path(td)/"ClockBind_preregistration_package"; root.mkdir()
        for role,p in src:
            dst=root/_safe_name(p); shutil.copy2(p,dst)
            entries.append({"role":role,"file":dst.name,"sha256":sha256_file(dst),"bytes":dst.stat().st_size})
        manifest={"clockbind_package":"preregistration","status":status,"created_utc":created,
                  "frozen_by":frozen_by if status=="FROZEN" else "","external_registration":False,
                  "note":note,"files":entries}
        (root/"PREREG_MANIFEST.json").write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding="utf-8")
        (root/"README.txt").write_text(
            "ClockBind preregistration snapshot\n\n"
            f"Status: {status}\nCreated UTC: {created}\n"
            + (f"Frozen by: {frozen_by}\n" if status=="FROZEN" else "")
            + "This package is a local immutable snapshot. It does NOT prove OSF or registry submission.\n"
            + "Verify every listed SHA-256 before relying on the snapshot.\n",encoding="utf-8")
        with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
            for p in sorted(root.rglob("*")):
                if p.is_file(): z.write(p,p.relative_to(root.parent))
    return out


def extract_docx_text(path: str | Path) -> str:
    from docx import Document
    doc=Document(str(path)); chunks=[p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            chunks.extend(c.text for c in row.cells)
    return "\n".join(chunks)


def consistency_audit(manuscript: str | Path, registry: str | Path) -> tuple[list[dict], dict]:
    """Check a manuscript against an explicit result/claim registry.

    Registry format: {"checks":[{"id":"N","expected":"150","forbidden":["n=90"],
    "required":true}, ...]}.  This is intentionally explicit and auditable; the
    tool does not guess which numerical values should appear in a manuscript.
    """
    mp=Path(manuscript); rp=Path(registry)
    if not mp.exists(): raise FileNotFoundError(mp)
    if not rp.exists(): raise FileNotFoundError(rp)
    if mp.suffix.lower()==".docx": text=extract_docx_text(mp)
    else: text=mp.read_text(encoding="utf-8",errors="replace")
    obj=json.loads(rp.read_text(encoding="utf-8")); checks=obj.get("checks",[])
    rows=[]
    for c in checks:
        expected=str(c.get("expected", "")); required=bool(c.get("required", True)); forbidden=[str(x) for x in c.get("forbidden",[])]
        found=expected in text if expected else False; hits=[x for x in forbidden if x in text]
        ok=(found or not required) and not hits
        rows.append({"id":c.get("id",""),"required":required,"expected_found":found,"forbidden_hits":len(hits),"status":"PASS" if ok else "FAIL"})
    summ={"checks":len(rows),"pass":sum(r["status"]=="PASS" for r in rows),"fail":sum(r["status"]=="FAIL" for r in rows),
          "manuscript_sha256":sha256_file(mp),"registry_sha256":sha256_file(rp)}
    return rows,summ


def build_submission_bundle(manuscript: str | Path, artifacts: list[str | Path], output: str | Path) -> Path:
    mp=Path(manuscript).expanduser()
    if not mp.exists(): raise FileNotFoundError(mp)
    paths=[mp]+[Path(x).expanduser() for x in artifacts]
    for p in paths:
        if not p.exists(): raise FileNotFoundError(p)
    out=Path(output).expanduser(); out.parent.mkdir(parents=True,exist_ok=True)
    created=dt.datetime.now(dt.timezone.utc).isoformat(); entries=[]
    with tempfile.TemporaryDirectory() as td:
        root=Path(td)/"ClockBind_submission_bundle"; root.mkdir()
        for p in paths:
            dst=root/_safe_name(p); shutil.copy2(p,dst); entries.append({"file":dst.name,"sha256":sha256_file(dst),"bytes":dst.stat().st_size})
        (root/"SUBMISSION_MANIFEST.json").write_text(json.dumps({"created_utc":created,"files":entries},indent=2),encoding="utf-8")
        with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
            for p in sorted(root.rglob("*")):
                if p.is_file(): z.write(p,p.relative_to(root.parent))
    return out
