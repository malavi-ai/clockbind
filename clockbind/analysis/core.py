"""Analysis framework shared by the Studio app, the CLI and syntax files.

An analysis is a function f(df, **params) -> Output, registered under a name.
A *syntax* entry is {"analysis": name, "params": {...}} — the same dict the
Studio shows under each result, so every point-and-click result can be re-run
from the command line:  clockbind syntax run --data file.xlsx --syntax syntax.json
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd

from .. import __version__

REGISTRY: dict[str, dict] = {}


def analysis(name: str, title: str, group: str, params: list[dict] | None = None, doc: str = ""):
    """Register an analysis. params describe the GUI form (see studio)."""
    def deco(fn: Callable):
        REGISTRY[name] = {"fn": fn, "title": title, "group": group, "params": params or [], "doc": doc or (fn.__doc__ or "").strip()}
        return fn
    return deco


@dataclass
class Block:
    kind: str                      # "heading" | "text" | "table" | "figure" | "note" | "methods" | "warning"
    content: object
    title: str = ""
    note: str = ""


@dataclass
class Output:
    analysis: str
    title: str
    params: dict
    blocks: list[Block] = field(default_factory=list)
    references: list[str] = field(default_factory=list)
    created: str = field(default_factory=lambda: _dt.datetime.now().isoformat(timespec="seconds"))
    data_sha256: str = ""
    n_used: int | None = None

    # builders
    def text(self, s): self.blocks.append(Block("text", s)); return self
    def warn(self, s): self.blocks.append(Block("warning", s)); return self
    def note(self, s): self.blocks.append(Block("note", s)); return self
    def methods(self, s): self.blocks.append(Block("methods", s)); return self
    def table(self, df: pd.DataFrame, title: str = "", note: str = ""):
        self.blocks.append(Block("table", df, title, note)); return self
    def figure(self, fig, title: str = "", note: str = ""):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
        import matplotlib.pyplot as plt
        plt.close(fig)
        self.blocks.append(Block("figure", buf.getvalue(), title, note)); return self
    def ref(self, *refs):
        for r in refs:
            if r not in self.references:
                self.references.append(r)
        return self

    @property
    def syntax(self) -> dict:
        return {"analysis": self.analysis, "params": self.params}


def data_hash(df: pd.DataFrame) -> str:
    return hashlib.sha256(pd.util.hash_pandas_object(df, index=True).values.tobytes()).hexdigest()


def run(name: str, df: pd.DataFrame, params: dict) -> Output:
    if name not in REGISTRY:
        raise KeyError(f"Unknown analysis '{name}'. Available: {sorted(REGISTRY)}")
    spec = REGISTRY[name]
    out: Output = spec["fn"](df, **params)
    out.analysis, out.params = name, dict(params)
    out.title = out.title or spec["title"]
    out.data_sha256 = data_hash(df)
    return out


# ----------------------------------------------------------------- helpers used by modules
def numeric(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(f"Variables not found: {missing}")
    return df[cols].apply(pd.to_numeric, errors="coerce")


def fmt_p(p: float) -> str:
    if p is None or (isinstance(p, float) and np.isnan(p)):
        return ""
    return "< .001" if p < 0.001 else f"{p:.3f}".lstrip("0")


def as_list(x) -> list[str]:
    if x is None:
        return []
    if isinstance(x, str):
        return [s.strip() for s in x.split(",") if s.strip()]
    return list(x)


# ----------------------------------------------------------------- export
def outputs_to_docx(outputs: list[Output], path: str | Path, title: str = "ClockBind output") -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.8)
    from .style import BRAND
    st = doc.styles["Normal"]
    st.font.name, st.font.size = BRAND.get("fonts", {}).get("word", "Calibri"), Pt(10.5)
    logo = BRAND.get("word_header_logo") or ""
    lp = Path(logo) if logo and Path(logo).is_absolute() else Path(__file__).resolve().parents[1] / "assets" / logo if logo else None
    if lp and lp.exists():
        doc.sections[0].header.paragraphs[0].add_run().add_picture(str(lp), height=Cm(1.0))
    if BRAND.get("word_footer"):
        fp = doc.sections[0].footer.paragraphs[0]; fr = fp.add_run(BRAND["word_footer"]); fr.font.size = Pt(8)
    h = doc.add_heading(title, 0)
    doc.add_paragraph(f"ClockBind {__version__} · generated {_dt.datetime.now():%Y-%m-%d %H:%M}").runs[0].font.color.rgb = RGBColor(0x66, 0x72, 0x81)
    refs = []
    for o in outputs:
        doc.add_heading(o.title, 1)
        meta = doc.add_paragraph()
        r = meta.add_run(f"Data sha256 {o.data_sha256[:16]} · {o.created}" + (f" · n = {o.n_used}" if o.n_used is not None else ""))
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor(0x66, 0x72, 0x81)
        for b in o.blocks:
            if b.kind == "table":
                if b.title:
                    p = doc.add_paragraph(); rr = p.add_run(b.title); rr.bold = True
                df = b.content.reset_index() if not isinstance(b.content.index, pd.RangeIndex) else b.content
                t = doc.add_table(rows=1, cols=len(df.columns))
                t.style = "Light Grid Accent 1"
                for j, c in enumerate(df.columns):
                    t.rows[0].cells[j].text = str(c)
                for _, row in df.iterrows():
                    cells = t.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = _fmt_cell(v)
                small = Pt(7.5) if len(df.columns) > 8 else Pt(9)
                for row_ in t.rows:
                    for cell in row_.cells:
                        for par in cell.paragraphs:
                            for run_ in par.runs:
                                run_.font.size = small
                if b.note:
                    p = doc.add_paragraph(); rr = p.add_run("Note. " + b.note); rr.italic = True; rr.font.size = Pt(9)
            elif b.kind == "figure":
                if b.title:
                    p = doc.add_paragraph(); rr = p.add_run(b.title); rr.bold = True
                doc.add_picture(io.BytesIO(b.content), width=Cm(15))
                if b.note:
                    p = doc.add_paragraph(); rr = p.add_run("Note. " + b.note); rr.italic = True; rr.font.size = Pt(9)
            elif b.kind == "methods":
                p = doc.add_paragraph(); rr = p.add_run("Methods text: "); rr.bold = True; p.add_run(str(b.content))
            elif b.kind == "warning":
                p = doc.add_paragraph(); rr = p.add_run("Caution: " + str(b.content)); rr.font.color.rgb = RGBColor(0xB3, 0x45, 0x2A)
            else:
                doc.add_paragraph(str(b.content))
        p = doc.add_paragraph(); rr = p.add_run("Syntax: " + json.dumps(o.syntax, ensure_ascii=False)); rr.font.size = Pt(8); rr.font.name = "Consolas"
        refs += [x for x in o.references if x not in refs]
    if refs:
        doc.add_heading("References", 1)
        for r_ in sorted(refs):
            doc.add_paragraph(r_)
    path = Path(path)
    doc.save(path)
    return path


def _fmt_cell(v) -> str:
    if isinstance(v, (float, np.floating)):
        if np.isnan(v):
            return ""
        return f"{v:.3f}" if abs(v) < 1e5 else f"{v:.3g}"
    return str(v)


def syntax_file(outputs: list[Output]) -> str:
    return json.dumps({"clockbind": __version__, "created": _dt.datetime.now().isoformat(timespec="seconds"),
                       "steps": [o.syntax for o in outputs]}, indent=2, ensure_ascii=False)
