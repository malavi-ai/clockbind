"""Full Bridge report: one command, every output.

Reads a Bridge master workbook (01 register, 02A search log, 02B reconciliation,
04 metadata, 05 level assessment, Episodes/Steps, 08 coder agreement,
09 change log), applies the registered gates and the binding rule locally, and
writes:

    Bridge_Report.pdf      one combined report (tables + charts)
    Bridge_Report.xlsx     every table, one sheet each
    Bridge_Report.docx     tables and figures ready to paste into a manuscript
    Bridge_Report.md       short text summary
    tables/*.csv           every table as CSV
    figures/*.png|svg      every chart (PNG 200 dpi + SVG)
    Bridge_Report_Bundle.zip  everything above + manifest

`chat_safe=True` keeps only aggregate counts and aggregate charts: no episode
codes, no per-episode rows, no change-log text. Use that mode for anything that
leaves this computer or is shown in a chat app.
"""
from __future__ import annotations

import datetime as _dt
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

INK, ACCENT, SECOND, ERROR, MUTED, LINE = "#14233A", "#A4772B", "#2A7A6D", "#B3452A", "#5C6978", "#D3DADB"
AMBER, LIGHT = "#C99A45", "#8FA3B8"
STATE_COLORS = {"Yes": SECOND, "No": ERROR, "PENDING": AMBER, "Held": LIGHT, "N/A": "#B9C3CC", "blank": "#ECEFF1"}
HUF_BANDS = [(0, 200_000, "< 200k"), (200_000, 1_000_000, "200k–1M"), (1_000_000, 5_000_000, "1M–5M"), (5_000_000, float("inf"), "≥ 5M")]


# ----------------------------------------------------------------- reading
def _sheet(xl: dict, *names):
    low = {k.strip().lower(): k for k in xl}
    for n in names:
        if n.lower() in low:
            return xl[low[n.lower()]]
    return None


def _read_all(path, header_rows: dict) -> dict:
    out = {}
    names = pd.ExcelFile(path).sheet_names
    for n in names:
        hr = header_rows.get(n, 2 if n[:2].isdigit() else 0)
        try:
            df = pd.read_excel(path, sheet_name=n, header=hr, dtype=object)
            df = df.dropna(how="all")
            df.columns = [str(c).strip() for c in df.columns]
            out[n] = df
        except Exception:
            pass
    return out


def _s(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v).strip()


def _col(df, *cands):
    if df is None:
        return None
    for c in cands:
        if c in df.columns:
            return c
    return None


# ----------------------------------------------------------------- charts
def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": LINE, "axes.labelcolor": INK,
                         "xtick.color": MUTED, "ytick.color": INK, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.titleweight": "bold", "axes.titlesize": 11, "axes.titlecolor": INK, "axes.titlelocation": "left"})
    return plt


class Figs:
    def __init__(self, folder: Path):
        self.folder = folder
        folder.mkdir(parents=True, exist_ok=True)
        self.items: list[tuple[str, Path]] = []

    def save(self, fig, name, title):
        plt = _plt()
        png, svg = self.folder / f"{name}.png", self.folder / f"{name}.svg"
        fig.savefig(png, dpi=200, bbox_inches="tight")
        fig.savefig(svg, bbox_inches="tight")
        plt.close(fig)
        self.items.append((title, png))
        return png


def chart_hbar(figs, name, title, labels, values, colors=None, note=""):
    plt = _plt()
    n = len(labels)
    fig, ax = plt.subplots(figsize=(7.6, 0.38 * n + 1.1))
    y = np.arange(n)[::-1]
    ax.barh(y, values, color=colors or INK, height=0.62)
    ax.set_yticks(y, [str(l)[:62] for l in labels])
    mx = max([v for v in values if v is not None] + [1])
    for yi, v in zip(y, values):
        ax.text(v + mx * 0.01, yi, f"{v}", va="center", fontsize=8.5, color=INK)
    ax.set_xlim(0, mx * 1.12)
    ax.xaxis.set_visible(False)
    ax.spines["bottom"].set_visible(False)
    ax.set_title(title)
    if note:
        fig.text(0.01, -0.02, note, fontsize=7.5, color=MUTED, ha="left")
    return figs.save(fig, name, title)


def chart_stacked_gates(figs, name, title, tab: pd.DataFrame):
    plt = _plt()
    states = [s for s in ["No", "PENDING", "Held", "N/A", "Yes"] if s in tab.columns]
    n = len(tab)
    fig, ax = plt.subplots(figsize=(7.6, 0.34 * n + 1.4))
    y = np.arange(n)[::-1]
    left = np.zeros(n)
    for s in states:
        v = tab[s].to_numpy(dtype=float)
        ax.barh(y, v, left=left, color=STATE_COLORS[s], height=0.62, label=s)
        left += v
    ax.set_yticks(y, [str(i)[:58] for i in tab.index])
    ax.set_title(title)
    ax.legend(ncol=len(states), loc="upper left", bbox_to_anchor=(0, -0.04), frameon=False, fontsize=8)
    ax.xaxis.set_visible(False)
    ax.spines["bottom"].set_visible(False)
    return figs.save(fig, name, title)


def chart_matrix(figs, name, title, mat: pd.DataFrame):
    plt = _plt()
    from matplotlib.colors import ListedColormap
    order = ["Yes", "No", "PENDING", "Held", "N/A", "blank"]
    code = {s: i for i, s in enumerate(order)}
    z = mat.apply(lambda col: col.map(lambda v: code.get(v if v in code else ("blank" if v == "" else "Held"), 5))).to_numpy()
    fig, ax = plt.subplots(figsize=(max(6, 0.36 * mat.shape[1] + 2.2), 0.36 * mat.shape[0] + 1.6))
    ax.imshow(z, cmap=ListedColormap([STATE_COLORS[s] for s in order]), vmin=0, vmax=len(order) - 1, aspect="auto")
    ax.set_yticks(range(mat.shape[0]), mat.index)
    ax.set_xticks(range(mat.shape[1]), mat.columns, rotation=90, fontsize=7.5)
    ax.set_xticks(np.arange(-.5, mat.shape[1], 1), minor=True)
    ax.set_yticks(np.arange(-.5, mat.shape[0], 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=1.2)
    ax.tick_params(which="minor", length=0)
    for s in ("bottom", "left"):
        ax.spines[s].set_visible(False)
    ax.set_title(title)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=STATE_COLORS[s], label=s) for s in order[:-1]], ncol=5, loc="upper left",
              bbox_to_anchor=(0, -0.22 if mat.shape[0] < 6 else -0.12), frameon=False, fontsize=8)
    return figs.save(fig, name, title)


def chart_slack(figs, name, title, res: pd.DataFrame):
    plt = _plt()
    d = res.dropna(subset=["decisive_slack_earliest_days", "decisive_slack_latest_days"], how="all")
    if d.empty:
        return None
    n = len(d)
    fig, ax = plt.subplots(figsize=(7.6, 0.5 * n + 1.4))
    y = np.arange(n)[::-1]
    for yi, (_, r) in zip(y, d.iterrows()):
        e, l = r["decisive_slack_earliest_days"], r["decisive_slack_latest_days"]
        e = np.nan if e is None else float(e)
        l = np.nan if l is None else float(l)
        stable = str(r.get("sign_stable", "")) == "Yes"
        ax.plot([e, l], [yi, yi], color=SECOND if stable else ERROR, lw=2.2, zorder=1)
        ax.scatter([e], [yi], color=INK, s=28, zorder=2, label="earliest bound" if yi == y[0] else None)
        ax.scatter([l], [yi], facecolor="white", edgecolor=INK, s=28, zorder=2, label="latest bound" if yi == y[0] else None)
        ax.text(np.nanmax([e, l]) + 0.4, yi, f"  {r.get('verdict', '')}", va="center", fontsize=8, color=INK)
    ax.axvline(0, color=MUTED, lw=1, ls="--")
    ax.set_yticks(y, d["episode"].astype(str))
    ax.set_xlabel("decisive slack (days) = tW − T*(k); ≥ 0 means R would have been on time")
    ax.set_title(title)
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, frameon=False, fontsize=8)
    ax.set_title(title, pad=22)
    return figs.save(fig, name, title)


def chart_gantt(figs, name, title, ep_row, steps: pd.DataFrame, relative=False):
    """Step bars from the time a step became ready (all predecessors finished, or the anchor) to its earliest
    documented finish; the hatched extension is the finish-date uncertainty (earliest → latest). tW is the dashed line.
    relative=True shows days from the anchor (day 0) instead of calendar dates (for publication)."""
    plt = _plt()

    def ts(v):
        try:
            return pd.to_datetime(v)
        except Exception:
            return pd.NaT
    st = steps.copy()
    st["fe"], st["fl"] = st["finish_earliest"].map(ts), st["finish_latest"].map(ts)
    st["fl"] = st["fl"].fillna(st["fe"])
    st = st.dropna(subset=["fe"])
    if st.empty:
        return None
    anchor = ts(ep_row.get("anchor_earliest"))
    if pd.isna(anchor):
        anchor = st["fe"].min()
    fin = dict(zip(st["step"].astype(str), st["fe"]))

    def ready(r):
        preds = [p.strip() for p in str(r.get("predecessors", "") or "").replace(",", ";").split(";") if p.strip() and p.strip().lower() != "nan"]
        vals = [fin[p] for p in preds if p in fin]
        return max(vals) if vals else anchor
    st["rd"] = st.apply(ready, axis=1)
    st = st.sort_values(["rd", "fe"])
    conv = (lambda t: (t - anchor).total_seconds() / 86400) if relative else (lambda t: t)
    clock_col = {"finance": ACCENT, "payment": AMBER, "fulfilment": SECOND, "logistics": "#3E6E9E", "operational-readiness": INK}
    fig, ax = plt.subplots(figsize=(7.8, 0.52 * len(st) + 1.9))
    y = np.arange(len(st))[::-1]
    span_days = max((st["fl"].max() - min(st["rd"].min(), anchor)).total_seconds() / 86400, 1.0)
    min_w = span_days * 0.012
    r_step = str(ep_row.get("r_step", "") or "").strip()
    for yi, (_, r) in zip(y, st.iterrows()):
        ck = str(r.get("clock", "")).strip().lower()
        c = clock_col.get(ck, MUTED)
        x0, x1, x2 = conv(r["rd"]), conv(r["fe"]), conv(r["fl"])
        dur = (r["fe"] - r["rd"]).total_seconds() / 86400
        w_days = max(dur, min_w)
        w = w_days if relative else pd.Timedelta(days=w_days)
        ax.barh(yi, w, left=x0, color=c, height=0.52, edgecolor="white", linewidth=0.6)
        end = x0 + w
        if r["fl"] > r["fe"]:
            ax.barh(yi, (x2 - x1), left=x1, color="white", edgecolor=c, hatch="////", height=0.52, linewidth=0.8)
            end = max(end, x2)
        lab = f"{dur:.0f} d" if dur >= 1 else ("<1 d" if dur > 0 else "0 d")
        ax.text(end, yi, f"  {lab}", va="center", fontsize=7.5, color=MUTED)
    twe, twl = ts(ep_row.get("tw_earliest")), ts(ep_row.get("tw_latest"))
    if pd.notna(twe):
        if pd.notna(twl) and twl != twe:
            ax.axvspan(conv(twe), conv(twl), color=ERROR, alpha=0.10, lw=0)
        ax.axvline(conv(twe), color=ERROR, lw=1.3, ls="--")
        ax.annotate("tW", (conv(twe), 1.0), xycoords=("data", "axes fraction"), xytext=(3, -2), textcoords="offset points",
                    color=ERROR, fontsize=8.5, fontweight="bold", va="top")
    ax.set_yticks(y, [f"{n}  (R)" if n == r_step else n for n in st["step"].astype(str)])
    for tl in ax.get_yticklabels():
        if tl.get_text().endswith("(R)"):
            tl.set_fontweight("bold")
    ax.grid(axis="x", color=LINE, lw=0.6, alpha=0.7)
    ax.set_axisbelow(True)
    if relative:
        ax.set_xlabel("days from the anchor (day 0)", color=MUTED)
    else:
        import matplotlib.dates as mdates
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d %b"))
        fig.autofmt_xdate()
    from matplotlib.patches import Patch
    used = [k for k in clock_col if k in set(st["clock"].astype(str).str.strip().str.lower())]
    handles = [Patch(color=clock_col[k], label=k) for k in used] + [Patch(facecolor="white", edgecolor=MUTED, hatch="////", label="finish-date uncertainty")]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, -0.16 if relative else -0.24), ncol=min(len(handles), 4), frameon=False, fontsize=7.5)
    ax.set_title(title)
    return figs.save(fig, name, title)


def chart_ladder(figs, name, title, labels, values, total=None, note=""):
    """Descending ladder: each bar is the number of episodes reaching that layer of evidence."""
    plt = _plt()
    n = len(labels)
    fig, ax = plt.subplots(figsize=(7.6, 0.62 * n + 1.2))
    y = np.arange(n)[::-1]
    mx = max([total or 0] + [v for v in values] + [1])
    shades = [INK, "#3E6E9E", SECOND, ACCENT, ERROR][:n] if n <= 5 else [INK] * n
    if total:
        ax.barh(y, [total] * n, color="#EEF1F3", height=0.6)
    ax.barh(y, values, color=shades, height=0.6)
    for yi, v in zip(y, values):
        lab = f"{v}" + (f" of {total}" if total else "")
        ax.text(v + mx * 0.012, yi, lab, va="center", fontsize=9, color=INK, fontweight="bold")
    ax.set_yticks(y, [str(l)[:58] for l in labels])
    ax.set_xlim(0, mx * 1.18)
    ax.xaxis.set_visible(False)
    ax.spines["bottom"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_title(title)
    if note:
        fig.text(0.01, -0.03, note, fontsize=7.5, color=MUTED, ha="left")
    return figs.save(fig, name, title)


# ----------------------------------------------------------------- statistics helpers
def agreement(x, y):
    x, y = list(x), list(y)
    n = len(x)
    if n == 0:
        return np.nan, np.nan, np.nan
    cats = sorted(set(x) | set(y))
    po = float(np.mean([a == b for a, b in zip(x, y)]))
    px = {c: np.mean([a == c for a in x]) for c in cats}
    py = {c: np.mean([b == c for b in y]) for c in cats}
    pe = sum(px[c] * py[c] for c in cats)
    kappa = (po - pe) / (1 - pe) if pe < 1 else np.nan
    q = len(cats)
    pi = {c: (px[c] + py[c]) / 2 for c in cats}
    pe1 = sum(p * (1 - p) for p in pi.values()) / (q - 1) if q > 1 else np.nan
    ac1 = (po - pe1) / (1 - pe1) if q > 1 and pe1 < 1 else np.nan
    return po, kappa, ac1


def claim_tier(level, verdict, sign, rival, tier, setter, l3e):
    """Claim ceiling (frozen v3.4). Positive binding verdicts get A/B/C by strength of evidence; Indeterminate or a
    sign-unstable result → X; Non-binding gets the negated form at the same tier (shown as 'A¬', 'B¬', 'C¬') and is
    never counted as a positive binding attribution; no verdict → PENDING."""
    level, verdict = _s(level), _s(verdict)
    if level == "Level 3":
        if not verdict:
            return "PENDING"
        if verdict == "Indeterminate" or _s(sign) == "No":
            return "X"
        neg = verdict == "Non-binding"
        t = "PENDING"
        if _s(rival) == "Not rejected":
            t = "C"
        elif _s(rival) == "Rejected" and _s(sign) == "Yes" and _s(l3e) == "Yes":
            if _s(tier) == "A" and _s(setter) != "Seller":
                t = "A"
            elif _s(tier) == "B" or _s(setter) == "Seller":
                t = "B"
        return (t + "¬") if (neg and t != "PENDING") else t
    if level.startswith("Level 2"):
        return "D"
    if level.startswith("Level 1"):
        return "E"
    return "—"


NEG_WORDING = {"A": "was not binding under the specified decision rule", "B": "appears not to have been binding",
               "C": "is consistent with a non-binding clock"}


def claim_wording(tier, verdict=""):
    if tier.endswith("¬"):
        return NEG_WORDING.get(tier[:-1], "")
    return WORDING.get(tier, "")


WORDING = {"A": "was binding under the specified decision rule", "B": "appears to have been binding",
           "C": "is consistent with a binding constraint", "D": "suggests (mechanism only)",
           "E": "a delay was observed (description only)", "X": "binding status cannot be determined",
           "PENDING": "no claim yet", "—": "no claim"}



NEUTRAL_CODE = __import__("re").compile(r"^(E\d{1,3}[a-z]?|[A-Z]{2}-M\d{1,3}[A-Z]?|EX-?\d+|ZZ\d+)$")
PAYMENT_WORDS = ("pay", "deposit", "instal", "transfer", "remit")


def consistency(eps, eid, first_rows, res, E, S, rec, l3_cols) -> pd.DataFrame:
    """Cross-checks between the coded workbook and the binding inputs. Each row: level, check, n, episodes."""
    rows = []

    def add(level, check, eps_list, note=""):
        eps_list = sorted({str(e) for e in eps_list if _s(e)})
        rows.append({"level": level if eps_list else "PASS", "check": check, "n": len(eps_list),
                     "episodes": ", ".join(eps_list), "what to do": note if eps_list else ""})

    lvl = dict(zip(eps[eid].astype(str), eps["level"].map(_s)))
    l3 = [e for e, l in lvl.items() if l == "Level 3"]
    fr = lambda e, k: _s(first_rows.loc[e].get(k)) if (len(first_rows) and e in first_rows.index) else ""
    ep_in_E = set(E["episode"].map(_s)) if E is not None and "episode" in E.columns else set()

    add("FAIL", "Level-3 episode missing from the Episodes/Steps sheets", [e for e in l3 if e not in ep_in_E],
        "Add its timeline (Episodes + Steps) or the binding verdict cannot be computed.")
    add("WARN", "Episodes/Steps timeline for an episode that is not Level 3", [e for e in ep_in_E if e and lvl.get(e) != "Level 3"],
        "Binding verdicts are reported for Level 3 only; check the level or remove the timeline.")

    # L3b coded Yes but ex-ante durations missing in Steps
    l3b = l3_cols.get("L3b")
    miss = []
    if S is not None and "expected_days" in S.columns and l3b:
        for e in sorted(set(l3) | {x for x in ep_in_E if x in lvl}):
            if _s(eps.loc[eps[eid].astype(str) == e, l3b].iloc[0]) == "Yes":
                st = S[S["episode"].map(_s) == e]
                if len(st) and pd.to_numeric(st["expected_days"], errors="coerce").isna().any():
                    miss.append(f"{e} ({int(pd.to_numeric(st['expected_days'], errors='coerce').isna().sum())}/{len(st)} steps blank)")
    add("FAIL", "L3b = Yes in 05, but expected_days blank in Steps", miss,
        "Fill expected_days from documents dated before the outcome, or code L3b = No (episode drops to Level 2).")

    # zero ex-ante durations
    zero = []
    if S is not None and "expected_days" in S.columns:
        z = S[pd.to_numeric(S["expected_days"], errors="coerce") == 0]
        zero = [f"{_s(r['episode'])}:{_s(r['step'])}" for _, r in z.iterrows()]
    add("WARN", "Ex-ante duration of exactly 0 days", zero, "Confirm a pre-outcome document states 'immediate/same day'; a 0 can decide the verdict.")

    if len(res):
        # verdict recorded in 05 vs engine
        mm = []
        for _, r in res.iterrows():
            e = _s(r["episode"]); rec_v = fr(e, "Verdict")
            if rec_v and rec_v != _s(r["verdict"]):
                mm.append(f"{e} (05: {rec_v} / engine: {_s(r['verdict'])})")
        add("FAIL", "Verdict recorded in 05 differs from the binding engine", mm,
            "Find out which input changed (Change Log); record the engine verdict with date, or correct the timeline.")
        add("WARN", "Level-3 episode with no verdict recorded in 05", [e for e in l3 if not fr(e, "Verdict")],
            "Record the verdict and the sign-stability result in 05.")
        edge = [_s(r["episode"]) for _, r in res.iterrows() if _s(r.get("fragile")) == "Yes"]
        add("WARN", "Fragile verdict: decisive slack within ±1 day of zero", edge,
            "Pre-specified robustness flag (does not change the verdict): report both bounds.")
        zr = [f"{_s(r['episode'])}:{f.split(':')[0].replace('step ', '')}" for _, r in res.iterrows()
              for f in _s(r.get("flags")).split("; ") if "treated as missing" in f]
        add("WARN", "Ex-ante duration 0 without zero_documented = Yes (treated as missing)", zr,
            "Set Steps.zero_documented = Yes only if a pre-outcome document states immediate/same-day; otherwise enter the documented duration.")
        add("WARN", "Finance actionability not computable", [_s(r["episode"]) for _, r in res.iterrows() if "not computable" in _s(r.get("finance_actionable")).lower()],
            "Needs ex-ante durations of the steps after the finance step.")
    # rival test for Level 3
    rv = [e for e in l3 if fr(e, "Supplier-delay rival") not in ("Rejected", "Not rejected")]
    add("FAIL", "Level-3 episode without a valid supplier-delay rival result", rv, "Enter Rejected or Not rejected (required for every Level-3 episode).")

    # clock plausibility
    if S is not None and {"step", "clock"} <= set(S.columns):
        pc = [f"{_s(r['episode'])}:{_s(r['step'])}" for _, r in S.iterrows()
              if any(w in _s(r["step"]).lower() for w in PAYMENT_WORDS) and _s(r["clock"]).lower() == "finance"]
        add("WARN", "Payment-type step coded on the finance clock", pc, "Check: payment execution belongs to the payment clock unless the step is financing becoming usable.")
        rr = [f"{_s(r['episode'])}:{_s(r['step'])}" for _, r in S.iterrows()
              if any(w in _s(r["step"]).lower() for w in ("operational", "install", "commission", "training", "opening", "go-live"))
              and _s(r["clock"]).lower() not in ("operational-readiness", "operational readiness")]
        add("WARN", "Readiness-type step not on the operational-readiness clock", rr, "Check the clock of this step.")

    # 02B link
    if rec is not None and "Decision" in rec.columns:
        retained = set(rec.loc[rec["Decision"].map(_s) == "Retained", "Candidate ID"].map(_s))
        add("FAIL", "Episode in 05 that is not Retained in 02B", [e for e in lvl if e not in retained], "Record its 02B decision, or remove the 05 row with a logged reason.")
        add("WARN", "Retained in 02B but missing from 05", [e for e in retained if e not in lvl], "Add the episode to 05.")
        add("WARN", "02B candidate without a decision", list(rec.loc[rec["Decision"].map(_s) == "", "Candidate ID"].map(_s)), "Retained / Merged / Not admitted, with a reason.")

    if first_rows is not None and len(first_rows):
        scol = next((c for c in SENS_COL_ALIASES if c in first_rows.columns), None)
        l3e = l3_cols.get("L3e")
        if scol and l3e:
            reached = [e for e, l in lvl.items() if l.startswith(("Level 2", "Level 3"))]
            need = [e for e in reached if e in first_rows.index and _s(first_rows.loc[e].get(l3e)) != "Yes" and sens_kind(first_rows.loc[e].get(scol)) in ("", "pending")]
            add("WARN", "L3e not met but sensitivity source type not coded (S1/S2 incomplete)", need,
                "Code the source type from the documents before the lock, or justify (e.g. the S1 upper bound is zero).")
    add("WARN", "Episode code does not look like a neutral pseudonym", [e for e in lvl if not NEUTRAL_CODE.match(e)],
        "Replace codes derived from names or companies (e.g. E20, E21 …); keep the key company-side.")
    return pd.DataFrame(rows)


SENS_COL = "L3e-S source type"
SENS_COL_ALIASES = ("L3e-S source type", "L3e-S customer-date source type")
SENS_VALUES = ["System-generated", "Standard transactional (manual)", "Narrative / retrospective", "External source", "PENDING"]


def sens_column(df):
    for c in SENS_COL_ALIASES:
        if c in df.columns:
            return c
    return None


def sens_kind(v) -> str:
    """Normalise a coded source type: 'system', 'manual', 'narrative', 'external', 'pending' or ''."""
    t = _s(v).lower()
    if not t:
        return ""
    if "pending" in t:
        return "pending"
    if "system" in t:
        return "system"
    if "standard" in t or "manual" in t or "transactional" in t:
        return "manual"
    if "narrative" in t or "retrospective" in t:
        return "narrative"
    if "external" in t:
        return "external"
    return "other"


def sensitivity(df, g, eid, eps_primary, ep_in_timeline, chat_safe=False):
    """Primary vs S1 vs S2 for the L3e independence requirement. Returns (summary table, per-episode table or None)."""
    from .plugins.screen import apply_gates
    l3e = next((c["column"] for s in g["stages"] for c in s["criteria"] if c["id"] == "L3e"), None)
    col = sens_column(df)
    if l3e is None or col is None:
        return None, None
    accept = {"S1": {"system", "manual"}, "S2": {"system"}}
    levels = {"Primary": eps_primary.set_index(eid)["level"].astype(str)}
    for name, ok in accept.items():
        d = df.copy()
        src = d[col].map(sens_kind)
        d[l3e] = [("Yes" if (_s(v) == "Yes" or st in ok) else v) for v, st in zip(d[l3e], src)]
        _, e2, _ = apply_gates(d, g)
        levels[name] = e2.set_index(eid)["level"].astype(str)
    lv = pd.DataFrame(levels)
    fails = eps_primary.set_index(eid).get("L3_failed", pd.Series(dtype=object)).map(_s)
    upper = int(sum(1 for v in fails if v and set(x.strip() for x in v.replace(";", ",").split(",") if x.strip()) == {"L3e"}))
    rows = []
    for name in ("Primary", "S1", "S2"):
        l3 = lv.index[lv[name] == "Level 3"]
        changed = lv.index[lv[name] != lv["Primary"]]
        rows.append({"analysis": name, "Level 3": int(len(l3)), "episodes changing level vs Primary": int(len(changed)),
                     "new Level-3 episodes without a timeline": int(sum(1 for e in l3 if e not in ep_in_timeline and lv.loc[e, "Primary"] != "Level 3")),
                     "L3e source type coded": int(df[col].map(sens_kind).isin(["system", "manual", "narrative", "external", "other"]).sum())})
    summ = pd.DataFrame(rows)
    summ.loc[len(summ)] = {"analysis": "Upper bound of S1 (fail L3e only)", "Level 3": int((lv["Primary"] == "Level 3").sum()) + upper,
                           "episodes changing level vs Primary": upper, "new Level-3 episodes without a timeline": None, "L3e source type coded": None}
    per = None if chat_safe else lv[(lv["S1"] != lv["Primary"]) | (lv["S2"] != lv["Primary"])].reset_index().rename(columns={"index": eid})
    return summ, per

# ----------------------------------------------------------------- main builder
def build(data: str, gates: str, out_root: str, sheet: str = "05_Level_Assessment", header_row: int = 2,
          chat_safe: bool = False, title: str = "Bridge — full report", run=None, relative_time: bool = False) -> dict:
    from .core.provenance import sha256_file
    from .pdfreport import build_pdf
    from .plugins.screen import apply_gates, check_frozen, funnel, flow_svg, load_data_sheet, load_gates, criterion_labels

    out = Path(out_root)
    tdir, fdir = out / "tables", out / "figures"
    tdir.mkdir(parents=True, exist_ok=True)
    figs = Figs(fdir)
    tables: dict[str, pd.DataFrame] = {}
    blocks: list = []
    notes: list[str] = []
    mode = "CHAT-SAFE (aggregate only)" if chat_safe else "FULL (local use only — contains pseudonymised episode codes)"

    g = load_gates(gates)
    h, frozen_ok = check_frozen(g, gates)
    xl = _read_all(data, {sheet: header_row, "Episodes": 0, "Steps": 0})

    # 1 status
    status = pd.DataFrame([
        ("Report created", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("Mode", mode),
        ("Workbook", Path(data).name), ("Workbook SHA-256", sha256_file(data)[:16]),
        ("Gates", f"{g.get('protocol_version', '?')} · sha256 {h[:16]}"),
        ("Gates frozen", "Yes" if g.get("frozen") else "No — results are a draft and must not be reported")], columns=["item", "value"])
    tables["00_Status"] = status

    # 2 workbook integrity
    try:
        from .workbook_check import check_workbook
        rep = check_workbook(data).frame()
        wsum = rep.groupby("level").size().reindex(["FAIL", "WARN", "PASS"]).fillna(0).astype(int).rename_axis("result").reset_index(name="checks")
        tables["01_Integrity_summary"] = wsum
        tables["01_Integrity_detail"] = rep[["level", "check", "where"]] if chat_safe else rep
    except Exception as e:
        notes.append(f"Workbook integrity check skipped: {e}")

    # 3 registers and pipeline
    pipe = []
    reg = _sheet(xl, "01_Source_Register")
    if reg is not None and _col(reg, "Source code"):
        reg = reg[reg["Source code"].map(_s) != ""]
        an = reg.apply(lambda r: "Yes" if all(_s(r.get(c)) == "Yes" for c in ("Unique?", "In archive window?", "Access cleared?")) else "No", axis=1)
        log = _sheet(xl, "02A_Search_Log")
        searched = set(log["Source code"].map(_s)) if log is not None and "Source code" in log.columns else set()
        pipe += [("Raw records registered", len(reg)), ("Unique records", int((reg.get("Unique?", pd.Series(dtype=object)).map(_s) == "Yes").sum())),
                 ("Analytic sources", int((an == "Yes").sum())),
                 ("Analytic sources searched with fixed terms", (lambda k: k if k > 0 else
                    "not documented at source level in 02A_Search_Log; historical hit counts not reconstructed")(
                    int(((an == "Yes") & reg["Source code"].map(_s).isin(searched)).sum())))]
        if "Source class" in reg.columns:
            sc = reg["Source class"].map(_s).replace("", "not recorded").value_counts()
            tables["03_Source_classes"] = sc.rename_axis("source class").reset_index(name="records")
    rec = _sheet(xl, "02B_Episode_Reconciliation")
    if rec is not None and _col(rec, "Candidate ID"):
        rec = rec[rec["Candidate ID"].map(_s) != ""]
        dec = rec.get("Decision", pd.Series(dtype=object)).map(_s)
        pipe += [("Episode candidates", len(rec)), ("Retained", int((dec == "Retained").sum())),
                 ("Merged", int((dec == "Merged").sum())), ("Not admitted", int((dec == "Not admitted").sum()))]

    # 4 screening engine
    df = load_data_sheet(data, sheet, header_row)
    entries, eps, probs = apply_gates(df, g)
    n_err = int((probs["level"] == "ERROR").sum()) if len(probs) else 0
    citable = bool(frozen_ok and g.get("frozen") and n_err == 0)
    status.loc[len(status)] = ["Screening checks", f"{n_err} errors, {int((probs['level'] == 'WARN').sum()) if len(probs) else 0} warnings"]
    status.loc[len(status)] = ["Citable", "Yes" if citable else "No"]
    fun = funnel(entries, eps, g)
    tables["02_Screening_funnel"] = fun
    if len(probs):
        if chat_safe:
            pc = [c for c in probs.columns if c.lower() not in ("where", "id", "episode", "episode id", "row", "entry")]
            tables["02_Screening_checks"] = probs.groupby(pc).size().reset_index(name="rows")
        else:
            tables["02_Screening_checks"] = probs
    try:
        (fdir / "flow_diagram.svg").write_text(flow_svg(fun, g), encoding="utf-8")
    except Exception:
        pass
    eid = "Episode ID" if "Episode ID" in eps.columns else eps.columns[0]
    lv = eps["level"].map(_s).replace("", "not assigned")
    order = ["Outside analytic archive", "Held at S0", "Below Level 1", "Level 1", "Level 2", "Level 3"]
    lvc = lv.value_counts()
    def _rank(x):
        for i, o in enumerate(order):
            if x == o or x.startswith(o + " "):
                return (i, x)
        return (99, x)
    lev = pd.DataFrame({"level": sorted(lvc.index, key=_rank)})
    lev["episodes"] = lev["level"].map(lvc).astype(int)
    tables["04_Levels"] = lev
    pipe += [("Episodes screened", len(eps))] + [(f"  {r.level}", int(r.episodes)) for r in lev.itertuples()]
    tables["02_Pipeline"] = pd.DataFrame(pipe, columns=["stage", "count"])

    # 5 gate breakdown (episode level)
    labels = criterion_labels(g)
    gate_rows, firsts = [], []
    gate_cols = []
    for stage in g["stages"]:
        for c in stage["criteria"]:
            col = c["column"]
            if col not in eps.columns:
                continue
            vals = eps[col].map(_s)
            norm = vals.map(lambda v: "Yes" if v in ("Yes", "COMPLETE") else ("No" if v == "No" else ("PENDING" if v == "PENDING" else ("N/A" if v == "N/A" else ("Held" if v else "blank")))))
            is_gate = c.get("gate", True) is not False
            row = {"gate": f"{c['id']} · {labels.get(c['id'], col)}"[:70], "id": c["id"], "is_gate": "gate" if is_gate else "flag/analysis"}
            for s in ("Yes", "No", "PENDING", "Held", "N/A", "blank"):
                row[s] = int((norm == s).sum())
            gate_rows.append(row)
            if is_gate and stage["id"] != "S0":
                gate_cols.append((c["id"], col, norm))
    gt = pd.DataFrame(gate_rows, columns=["gate", "id", "is_gate", "Yes", "No", "PENDING", "Held", "N/A", "blank"])
    tables["05_Gate_breakdown"] = gt
    for st_id in ("L1", "L2", "L3"):
        k = f"{st_id}_first_failed"
        if k in eps.columns:
            v = eps[k].map(_s)
            for gid, n in v[v != ""].value_counts().items():
                firsts.append((st_id, gid, int(n)))
    tables["05_First_failed_gate"] = pd.DataFrame(firsts, columns=["stage", "first failed gate", "episodes"])

    # 7 binding
    res = pd.DataFrame()
    E, S = _sheet(xl, "Episodes"), _sheet(xl, "Steps")
    if E is not None and S is not None and "episode" in E.columns and E["episode"].map(_s).ne("").any():
        from .binding import evaluate_all
        E = E[E["episode"].map(_s) != ""]
        S = S[S["episode"].map(_s) != ""]
        res = evaluate_all(E, S)
        vsum = res["verdict"].value_counts().rename_axis("verdict").reset_index(name="episodes")
        tables["07_Binding_verdicts"] = vsum
        stab = res["sign_stable"].map(_s).replace("", "not computed").value_counts().rename_axis("sign stable").reset_index(name="episodes")
        tables["07_Sign_stability"] = stab
        fa = res["finance_actionable"].map(_s).replace("", "not computed").value_counts().rename_axis("finance actionable").reset_index(name="episodes")
        tables["07_Finance_actionability"] = fa
        for col, key, lab in (("indeterminate_reason", "07_Indeterminate_reason", "indeterminate reason"),
                              ("fragile", "07_Fragile", "fragile (|slack| ≤ 1 day)")):
            if col in res.columns and res[col].map(_s).ne("").any():
                tables[key] = res[col].map(_s).replace("", None).dropna().value_counts().rename_axis(lab).reset_index(name="episodes")
        if "expected_days" in S.columns and "clock" in S.columns:
            sd = S.assign(d=pd.to_numeric(S["expected_days"], errors="coerce")).dropna(subset=["d"])
            if len(sd):
                tables["07_Exante_days_by_clock"] = sd.groupby("clock")["d"].agg(steps="count", median="median", minimum="min", maximum="max").round(2).reset_index()
        num = [c for c in res.columns if c.endswith("_days")]
        if num:
            tables["07_Slack_descriptives"] = res[num].apply(pd.to_numeric, errors="coerce").describe().T[["count", "mean", "50%", "min", "max"]].round(2).rename(columns={"50%": "median"}).reset_index(names="measure")
        if not chat_safe:
            tables["07_Binding_per_episode"] = res

    # 6 claim ladder
    first_rows = entries.groupby(eid, sort=False).first() if eid in entries.columns else pd.DataFrame()
    claims = []
    eng = res.assign(episode=res["episode"].map(_s)).drop_duplicates("episode").set_index("episode") if len(res) else None
    l3e = next((c["column"] for s in g["stages"] for c in s["criteria"] if c["id"] == "L3e"), None)
    for _, r in eps.iterrows():
        fr = first_rows.loc[r[eid]] if r[eid] in first_rows.index else {}
        gv = (lambda k: fr.get(k) if hasattr(fr, "get") else None)
        verdict, sign, vsrc = _s(gv("Verdict")), _s(gv("Sign-stable?")), "05"
        e_ = _s(r[eid])
        if not verdict and _s(r["level"]) == "Level 3" and eng is not None and e_ in eng.index:
            # registered rule unchanged: the verdict is the frozen engine's result; 05 only records it
            verdict, sign, vsrc = _s(eng.loc[e_, "verdict"]), _s(eng.loc[e_, "sign_stable"]), "engine"
        t = claim_tier(r["level"], verdict, sign, gv("Supplier-delay rival"), gv("Window tier (A/B/C)"), gv("Window setter"), r.get(l3e) if l3e else "")
        ir = _s(eng.loc[e_, "indeterminate_reason"]) if (eng is not None and e_ in eng.index and "indeterminate_reason" in eng.columns and verdict == "Indeterminate") else ""
        claims.append({"episode": r[eid], "level": r["level"], "verdict": verdict, "verdict source": vsrc if verdict else "",
                       "claim tier": t, "permitted wording": claim_wording(t, verdict) + (f" ({ir})" if ir else "")})
    cl = pd.DataFrame(claims)
    if len(cl):
        tables["06_Claim_tiers"] = cl.groupby(["claim tier", "permitted wording"]).size().reset_index(name="episodes")
        if (cl["verdict source"] == "engine").any():
            notes.append(f"Claim tiers: {int((cl['verdict source'] == 'engine').sum())} Level-3 verdict(s) not recorded in 05 were taken from the frozen binding engine.")
        if not chat_safe:
            tables["06_Claims_per_episode"] = cl
    if not chat_safe and len(eps):
        keep = [eid, "level", "final_status"] + [f"{s}_first_failed" for s in ("L1", "L2", "L3") if f"{s}_first_failed" in eps.columns] + [f"{s}_first_held" for s in ("L1", "L2", "L3") if f"{s}_first_held" in eps.columns]
        tables["04_Episodes"] = eps[keep]

    # 7b consistency checks
    l3_cols = {c["id"]: c["column"] for st_ in g["stages"] for c in st_["criteria"] if c["column"] in eps.columns}
    cons = consistency(eps, eid, first_rows, res, E, S, rec if rec is not None and "Candidate ID" in (rec.columns if rec is not None else []) else None, l3_cols)
    from . import descriptive as _d
    dtabs, dchecks, dpresent = _d.evaluate(eps, first_rows, eid, g)
    if dchecks:
        extra = [{"level": lv_ if el else "PASS", "check": ck, "n": len(set(el)), "episodes": ", ".join(sorted(set(el))),
                  "what to do": nt if el else ""} for lv_, ck, el, nt in dchecks]
        cons = pd.concat([cons, pd.DataFrame(extra)], ignore_index=True)
    tables.update(dtabs)
    if not dpresent:
        notes.append("Descriptive fields (gap reasons, stop point, buyer stage) not found in the data sheet.")
    lad = _d.temporal_ladder(eps, g)
    if lad is not None:
        tables["05_Temporal_ladder"] = lad
    fl = _d.finance_ladder(eps, g, res if len(res) else None)
    if fl is not None:
        tables["07_Finance_ladder"] = fl
    if chat_safe:
        cons = cons.drop(columns=["episodes"])
    tables["11_Consistency_checks"] = cons
    nf, nw = int((cons["level"] == "FAIL").sum()), int((cons["level"] == "WARN").sum())
    status.loc[len(status)] = ["Consistency checks", f"{nf} FAIL, {nw} WARN (see section 0)"]
    if nf:
        citable = False
        status.loc[status["item"] == "Citable", "value"] = "No"

    # 7c sensitivity of the L3e independence requirement (pre-specified S1/S2)
    ep_tl = set(E["episode"].map(_s)) if E is not None and "episode" in E.columns else set()
    try:
        ssum, sper = sensitivity(df, g, eid, eps, ep_tl, chat_safe)
    except Exception as ex:  # never block the report
        ssum, sper = None, None
        notes.append(f"Sensitivity analysis skipped: {ex}")
    if ssum is not None:
        tables["12_Sensitivity_L3e"] = ssum
        if sper is not None and len(sper):
            tables["12_Sensitivity_changed_episodes"] = sper
    else:
        notes.append(f"Sensitivity S1/S2 not run: no column '{SENS_COL}' in {sheet}.")

    # 8 money bands (aggregate) and author side
    md_ = _sheet(xl, "04_Source_Metadata")
    if md_ is not None and "Source code" in md_.columns:
        md_ = md_[md_["Source code"].map(_s) != ""]
        if "Author side" in md_.columns:
            tables["08_Author_side"] = md_["Author side"].map(_s).replace("", "not recorded").value_counts().rename_axis("author side").reset_index(name="documents")
        amt, cur, rate = _col(md_, "Commitment amount (company-side)"), _col(md_, "Currency"), _col(md_, "HUF per 1 unit (ECB)")
        if amt and cur:
            a = pd.to_numeric(md_[amt], errors="coerce")
            r_ = pd.to_numeric(md_[rate], errors="coerce") if rate else pd.Series(np.nan, index=md_.index)
            huf = np.where(md_[cur].map(_s) == "HUF", a, a * r_)
            huf = pd.Series(huf, index=md_.index).dropna()
            if len(huf):
                band = huf.map(lambda v: next(lbl for lo, hi, lbl in HUF_BANDS if lo <= v < hi))
                tables["08_Commitment_bands_HUF"] = band.value_counts().reindex([b[2] for b in HUF_BANDS]).fillna(0).astype(int).rename_axis("HUF band").reset_index(name="documents")

    # 9 coder agreement
    ag = _sheet(xl, "08_Coder_Agreement")
    if ag is not None and {"Author code", "Coder code"} <= set(ag.columns):
        ag = ag[(ag["Author code"].map(_s) != "") & (ag["Coder code"].map(_s) != "")]
        if len(ag):
            po, k, a1 = agreement(ag["Author code"].map(_s), ag["Coder code"].map(_s))
            prev = float((ag["Author code"].map(_s) == "Yes").mean())
            primary = "AC1" if (prev > 0.8 or prev < 0.2) else "kappa"
            tables["09_Agreement"] = pd.DataFrame([("Items compared", len(ag)), ("Raw agreement", round(po, 3)), ("Cohen's kappa", round(k, 3)),
                                                   ("Gwet's AC1", round(a1, 3)), ("Author 'Yes' prevalence", round(prev, 3)),
                                                   ("Primary statistic (registered rule)", primary)], columns=["measure", "value"])

    # 10 change log
    cg = _sheet(xl, "09_Change_Log")
    if cg is not None and "Type" in cg.columns:
        cg = cg[cg["Type"].map(_s) != ""]
        tables["10_Change_log_types"] = cg["Type"].map(_s).value_counts().rename_axis("type").reset_index(name="entries")
        if not chat_safe and "Entry" in cg.columns:
            ud = cg[cg["Type"].map(_s) == "USER DECISION"]
            if len(ud):
                tables["10_User_decisions"] = ud[[c for c in ("Date", "Entry", "By") if c in ud.columns]]

    # ------------------------------------------------------------- figures
    p = tables["02_Pipeline"]
    p = p[p["count"].map(lambda v: isinstance(v, (int, np.integer)))]
    chart_hbar(figs, "fig01_pipeline", "From archive records to screened levels", p["stage"].str.strip().tolist(), p["count"].tolist(),
               colors=[ACCENT if "Level 3" in s else SECOND if "Level 2" in s else "#5A7696" if "Level 1" in s else INK for s in p["stage"]],
               note="Counts describe this archive only; no prevalence claim.")
    chart_hbar(figs, "fig02_levels", "Episodes by evidence level", lev["level"].tolist(), lev["episodes"].tolist(),
               colors=[ACCENT if l == "Level 3" else SECOND if l == "Level 2" else "#5A7696" if l == "Level 1" else MUTED for l in lev["level"]])
    gg = gt[gt["is_gate"] == "gate"].set_index("gate")
    if len(gg):
        chart_stacked_gates(figs, "fig03_gates", "Where attribution breaks down — answers per gate", gg[["Yes", "No", "PENDING", "Held", "N/A"]])
    if len(cl):
        ctab = tables["06_Claim_tiers"].copy()
        order = {t: i for i, t in enumerate(["A", "B", "C", "A¬", "B¬", "C¬", "D", "E", "X", "PENDING", "—"])}
        ctab = ctab.sort_values("claim tier", key=lambda c: c.map(lambda t: order.get(t, 99)))
        chart_hbar(figs, "fig04_claims", "Claim ceiling reached", [f"{t} · {w}" for t, w in zip(ctab["claim tier"], ctab["permitted wording"])],
                   ctab["episodes"].astype(int).tolist(),
                   colors=[ACCENT if t in ("A", "B", "C") else LIGHT if t.endswith("¬") else SECOND if t == "D" else "#5A7696" if t == "E" else ERROR if t == "X" else MUTED for t in ctab["claim tier"]])
    if not chat_safe and gate_cols and len(eps):
        mat = pd.DataFrame({gid: norm.values for gid, _, norm in gate_cols}, index=eps[eid].astype(str).values)
        chart_matrix(figs, "fig05_evidence_matrix", "Evidence matrix — episodes × gates", mat)
    if len(res):
        if not chat_safe:
            chart_slack(figs, "fig06_slack", "Sign stability — decisive slack at earliest and latest bounds", res)
            for _, er in E.iterrows():
                eid_ = _s(er["episode"])
                chart_gantt(figs, f"fig07_timeline_{eid_}", f"Timeline {eid_} — steps, finish-date bounds and tW" + (" (days from anchor)" if relative_time else ""),
                            er, S[S["episode"].map(_s) == eid_], relative=relative_time)
        v = tables["07_Binding_verdicts"]
        chart_hbar(figs, "fig08_verdicts", "Binding verdicts (Level 3)", v["verdict"].tolist(), v["episodes"].tolist(), colors=ACCENT)
    if "05_Temporal_ladder" in tables:
        tl = tables["05_Temporal_ladder"]
        chart_ladder(figs, "fig10_temporal_ladder", "Temporal evidence ladder — realised time vs ex-ante time",
                     [st_.split(" · ", 1)[-1] for st_ in tl["step"]], tl["Yes"].astype(int).tolist(), total=int(tl["of analysable"].iloc[0]),
                     note="Analysable episodes documenting each layer of temporal evidence.")
    if "07_Finance_ladder" in tables:
        fl_ = tables["07_Finance_ladder"]
        fl_ = fl_[~fl_.iloc[:, 0].str.contains("not computable|of which", regex=True)]
        chart_ladder(figs, "fig11_finance_ladder", "From observed to binding financing",
                     [x.split(":")[0] for x in fl_.iloc[:, 0]], [int(v) for v in fl_["episodes"]],
                     note="Observed = financing route evaluated; actionable and binding from the frozen counterfactual engine.")
    if "03_Source_classes" in tables:
        t = tables["03_Source_classes"]
        chart_hbar(figs, "fig09_source_classes", "Registered records by source class", t["source class"].tolist(), t["records"].tolist(), colors=MUTED)

    # ------------------------------------------------------------- chat-safe safeguards
    share_note = None
    if chat_safe:
        from . import descriptive as _d
        manifest = _d.suppress_small_cells(tables)
        leaks = _d.leaked_codes({k: v for k, v in tables.items() if k != "00_Status"}, eps[eid].tolist())
        safe = not leaks
        status.loc[len(status)] = ["Share status", ("SAFE TO SHARE (aggregate only; counts 1–2 shown as <3 in source/descriptive tables)"
                                                    if safe else "NOT SAFE — episode codes found in tables; do not share")]
        lines = ["ClockBind chat-safe manifest", f"Workbook: {Path(data).name}  SHA-256 {sha256_file(data)}",
                 f"Generated: {_dt.datetime.now().isoformat(timespec='seconds')}",
                 "Removed: episode codes, per-episode tables, timelines, evidence matrix, dates, amounts.",
                 f"Small cells suppressed (<{_d.SMALL_CELL}): {len(manifest)}"] + [f"  - {k} · {c} · {r}" for k, c, r in manifest]
        lines.append("RESULT: SAFE TO SHARE" if safe else "RESULT: NOT SAFE — codes found: " + ", ".join(leaks))
        (out / ("SAFE_TO_SHARE.txt" if safe else "NOT_SAFE.txt")).write_text("\n".join(lines) + "\n", encoding="utf-8")
        share_note = "SAFE TO SHARE — see SAFE_TO_SHARE.txt" if safe else "NOT SAFE — see NOT_SAFE.txt"

    # ------------------------------------------------------------- writers
    for k, t in tables.items():
        t.to_csv(tdir / f"{k}.csv", index=False)
    xlsx = out / "Bridge_Report.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        for k, t in tables.items():
            t.to_excel(xw, sheet_name=k[:31], index=False)
    _style_xlsx(xlsx)

    figmap = {Path(pth).stem: (ttl, pth) for ttl, pth in figs.items}

    def fig(key):
        return [("image", str(figmap[key][1]))] if key in figmap else []

    blocks += [("kv", status.values.tolist()),
               ("note", "Descriptive of this archive under the registered rules. Analytic generalisation only: no prevalence claims. "
                        "Levels and verdicts are computed from the researcher's coding; the tool does not code cases.")]
    if not citable:
        blocks.append(("note", "NOT CITABLE: gates not frozen and/or screening errors. Use for checking only."))
    if share_note:
        blocks.append(("note", share_note))
    cview = tables["11_Consistency_checks"]
    blocks += [("h", "0 · Consistency checks (coding vs binding inputs)"),
               ("table", cview[cview["level"] != "PASS"] if (cview["level"] != "PASS").any() else cview),
               ("note", f"{int((cview['level'] == 'PASS').sum())} further checks passed. FAIL items must be fixed before the data lock.")]
    if "01_Integrity_summary" in tables:
        blocks += [("h", "1 · Workbook integrity"), ("table", tables["01_Integrity_summary"]), ("table", tables["01_Integrity_detail"].head(40))]
    blocks += [("h", "2 · Pipeline"), ("table", tables["02_Pipeline"])] + fig("fig01_pipeline")
    blocks += [("h", "3 · Screening funnel (registered gates)"), ("table", fun)]
    if "03_Source_classes" in tables:
        blocks += [("h", "4 · Source mix"), ("table", tables["03_Source_classes"])] + fig("fig09_source_classes")
        for k in ("08_Author_side", "08_Commitment_bands_HUF"):
            if k in tables:
                blocks.append(("table", tables[k]))
    blocks += [("h", "5 · Evidence levels"), ("table", lev)] + fig("fig02_levels")
    blocks += [("h", "6 · Where attribution breaks down"), ("table", gt.drop(columns=["id"]))] + fig("fig03_gates")
    if len(tables["05_First_failed_gate"]):
        blocks += [("p", "First failed gate per episode (the gate that stopped it):"), ("table", tables["05_First_failed_gate"])]
    if "05_Temporal_ladder" in tables:
        blocks += [("h", "6b · Temporal evidence ladder (analysable episodes)"), ("table", tables["05_Temporal_ladder"])] + fig("fig10_temporal_ladder") + [
                   ("note", "Realised time vs ex-ante time: how many episodes document each layer of temporal evidence.")]
    dkeys = [k for k in tables if k.startswith("13_")]
    if dkeys:
        blocks += [("h", "6c · Descriptive fields (never change a gate, level or verdict)")]
        for k in dkeys:
            blocks += [("p", k[3:].replace("_", " ") + ":"), ("table", tables[k])]
    blocks += fig("fig05_evidence_matrix")
    if "06_Claim_tiers" in tables:
        blocks += [("h", "7 · Claim ladder"), ("table", tables["06_Claim_tiers"])] + fig("fig04_claims")
        if "06_Claims_per_episode" in tables:
            blocks.append(("table", tables["06_Claims_per_episode"]))
    if len(res):
        blocks += [("h", "8 · Which clock binds? (Level 3)"), ("table", tables["07_Binding_verdicts"]), ("table", tables["07_Sign_stability"]),
                   ("table", tables["07_Finance_actionability"])] + fig("fig08_verdicts") + fig("fig11_finance_ladder") + fig("fig06_slack")
        for k in ("07_Finance_ladder", "07_Indeterminate_reason", "07_Fragile", "07_Exante_days_by_clock", "07_Slack_descriptives"):
            if k in tables:
                blocks.append(("table", tables[k]))
        if "07_Binding_per_episode" in tables:
            cols = [c for c in ["episode", "verdict", "indeterminate_reason", "binding_clocks", "sign_stable", "fragile", "finance_actionable", "reason"] if c in res.columns]
            blocks.append(("table", res[cols]))
        for key in sorted(k for k in figmap if k.startswith("fig07_")):
            blocks += fig(key)
    else:
        blocks += [("h", "8 · Which clock binds?"), ("p", "No Level-3 timelines in the Episodes/Steps sheets yet. Fill them after the freeze.")]
    if "12_Sensitivity_L3e" in tables:
        blocks += [("h", "8b · Sensitivity of the L3e independence requirement"), ("table", tables["12_Sensitivity_L3e"]),
                   ("note", "Primary = frozen gates. S1 also accepts contemporaneous routine transactional records of the firm; S2 only system-generated time-stamped records. Primary classifications are unchanged by these analyses.")]
        if "12_Sensitivity_changed_episodes" in tables:
            blocks.append(("table", tables["12_Sensitivity_changed_episodes"]))
    if "09_Agreement" in tables:
        blocks += [("h", "9 · Independent coder agreement"), ("table", tables["09_Agreement"])]
    if "10_Change_log_types" in tables:
        blocks += [("h", "10 · Change log"), ("table", tables["10_Change_log_types"])]
        if "10_User_decisions" in tables:
            blocks.append(("table", tables["10_User_decisions"]))
    for n_ in notes:
        blocks.append(("note", n_))

    pdf = out / "Bridge_Report.pdf"
    build_pdf(pdf, title, blocks, subtitle=f"{Path(data).name} · {mode}", landscape_pages=True)
    docx = out / "Bridge_Report.docx"
    _docx(docx, title, mode, tables, figs.items)
    md = out / "Bridge_Report.md"
    md.write_text(_markdown(title, status, tables, citable), encoding="utf-8")
    bundle = out / "Bridge_Report_Bundle.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(out.rglob("*")):
            if f.is_file() and f != bundle:
                z.write(f, f.relative_to(out))
    files = [pdf, xlsx, docx, md, bundle] + sorted(tdir.glob("*.csv")) + sorted(fdir.glob("*"))
    if run is not None:
        for f in files:
            run.add_output(f)
        run.note("mode", mode)
        run.note("citable", citable)
    return {"pdf": pdf, "xlsx": xlsx, "docx": docx, "md": md, "bundle": bundle, "tables": tables, "citable": citable, "summary": md.read_text(encoding="utf-8")}


def _style_xlsx(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            c.fill = PatternFill("solid", fgColor=INK[1:])
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=9)
        for col in ws.columns:
            w = max(len(str(c.value or "")) for c in col[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
        ws.freeze_panes = "A2"
    wb.save(path)


def _docx(path, title, mode, tables, figures):
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9.5)
    doc.add_heading(title, 0)
    doc.add_paragraph(f"{_dt.date.today().isoformat()} · {mode}")
    doc.add_paragraph("Descriptive of this archive under the registered rules; analytic generalisation only.")
    for k, t in tables.items():
        if len(t) > 60:
            continue
        doc.add_heading(k.split("_", 1)[1].replace("_", " "), 2)
        tb = doc.add_table(rows=1, cols=len(t.columns))
        tb.style = "Light Grid Accent 1"
        for i, c in enumerate(t.columns):
            tb.rows[0].cells[i].text = str(c)
        for _, r in t.iterrows():
            cells = tb.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = _s(v)
    doc.add_page_break()
    doc.add_heading("Figures", 1)
    for i, (ttl, p) in enumerate(figures, 1):
        doc.add_paragraph(f"Figure {i}. {ttl}").runs[0].bold = True
        doc.add_picture(str(p), width=Cm(16))
    doc.save(path)


def _markdown(title, status, tables, citable):
    lines = [f"# {title}", ""] + [f"- **{a}:** {b}" for a, b in status.values.tolist()] + [""]
    for k in ("11_Consistency_checks", "12_Sensitivity_L3e", "02_Pipeline", "04_Levels", "05_First_failed_gate", "05_Temporal_ladder", "06_Claim_tiers", "07_Binding_verdicts", "07_Finance_ladder", "09_Agreement"):
        if k in tables and len(tables[k]):
            lines += [f"## {k.split('_', 1)[1].replace('_', ' ')}", "", tables[k].to_markdown(index=False), ""]
    if not citable:
        lines.append("> NOT CITABLE — gates not frozen and/or screening errors.")
    return "\n".join(lines) + "\n"
