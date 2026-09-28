"""Case screening with frozen, versioned gates (Bridge study).

The gates live in a JSON config, not in code. The tool never takes a target
count: it applies the gates to every row and reports whatever numbers result.

Commands
  clockbind screen template  --gates G.json                    blank screening workbook
  clockbind screen freeze    --gates G.json --by NAME          lock gates; append to FREEZE_REGISTER.jsonl
  clockbind screen run       --data S.xlsx --gates G.json      apply gates, levels, funnel, flow diagram, checks
  clockbind screen compare   --data S.xlsx --gates A.json --gates-b B.json   who changes status/level
  clockbind screen agreement --data AUTHOR.xlsx --coder CODER.xlsx --gates G.json   agreement per criterion

Integrity model (v0.3)
  * Every coded value is validated against the value sets in the gates; anything
    else is an ERROR and is treated as Hold (never silently Fail or Pass).
  * IDs are trimmed; IDs that differ only in case/spacing are merged and flagged.
    Duplicate entry IDs stop the run.
  * Conflicting values inside one episode are Hold for every criterion type.
  * Criteria with "gate": false are recorded as flags and never exclude.
  * Freeze writes the gates hash to an append-only register next to the gates
    file; a protocol_version can be frozen once. The register is local, so for an
    external timestamp upload gates + register to OSF/Zenodo (see guide).
  * Every run is appended to RUN_LEDGER.jsonl (data hash, gates hash, counts).
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import sys
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd

from ..core.plugin import Plugin
from ..core.provenance import RunLog, sha256_file

PASS, FAIL, HOLD, NA = "Pass", "Fail", "Hold", "n/a"
CONFLICT = "CONFLICT"


# ----------------------------------------------------------------- config
def load_gates(path) -> dict:
    with open(path, encoding="utf-8") as f:
        g = json.load(f)
    for k in ("protocol_version", "id_column", "episode_column", "stages"):
        if k not in g:
            raise SystemExit(f"Gates file lacks '{k}'.")
    return g


def gates_hash(g: dict) -> str:
    core = {k: v for k, v in g.items() if k != "sha256"}
    return hashlib.sha256(json.dumps(core, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def register_path(gates_path) -> Path:
    return Path(gates_path).resolve().parent / "FREEZE_REGISTER.jsonl"


def read_register(gates_path) -> list[dict]:
    p = register_path(gates_path)
    if not p.exists():
        return []
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]


def check_frozen(g: dict, gates_path, problems: list | None = None, run=None) -> tuple[str, bool]:
    """Returns (hash, citable). Frozen gates must match their register entry."""
    h = gates_hash(g)
    if not g.get("frozen"):
        msg = f"Gates {g['protocol_version']} are NOT frozen: results are a draft and must not be reported."
        if run:
            run.warn(msg)
        return h, False
    if g.get("sha256") != h:
        raise SystemExit("Gates were edited after freezing (hash mismatch). Save them as a new protocol_version and freeze again.")
    reg = [r for r in read_register(gates_path) if r["protocol_version"] == g["protocol_version"]]
    if not reg:
        raise SystemExit(f"Gates claim to be frozen but {register_path(gates_path).name} has no entry for {g['protocol_version']}.")
    if reg[0]["sha256"] != h:
        raise SystemExit(f"Gates hash differs from the first registered freeze of {g['protocol_version']} ({reg[0]['frozen_on']}). Re-freezing a version is not allowed.")
    return h, True


def _paths(stage) -> list[list[dict]]:
    if "paths" in stage:
        return [p["criteria"] for p in stage["paths"]]
    return [stage["criteria"]]


def _path_names(stage) -> list[str]:
    if "paths" in stage:
        return [p["name"] for p in stage["paths"]]
    return [stage["name"]]


def _cols(c) -> list[str]:
    return c["columns"] if c.get("type") == "any_of" else [c["column"]]


def criterion_columns(g: dict, unit: str | None = None) -> list[str]:
    cols = []
    for st in g["stages"]:
        if unit and st.get("unit", "episode") != unit:
            continue
        for crits in _paths(st):
            for c in crits:
                cols += _cols(c)
    return list(dict.fromkeys(cols))


def _vals(g):
    v = g.get("values", {})
    return v.get("pass", ["Y"]), v.get("fail", ["N"]), v.get("unknown", ["Unknown"])


def column_specs(g) -> dict:
    """column -> (canonical allowed values, criterion) used for validation and AC1 scales."""
    yes, no, unk = _vals(g)
    spec = {}
    for st in g["stages"]:
        for crits in _paths(st):
            for c in crits:
                t = c.get("type", "yn")
                if t in ("yn", "any_of"):
                    allowed = list(yes) + list(no) + list(unk)
                elif t == "in":
                    allowed = list(dict.fromkeys(list(c.get("scale", c["allowed"])) + list(unk)))
                else:
                    allowed = None  # numeric
                for col in _cols(c):
                    if col in spec and spec[col][0] and allowed:
                        spec[col] = (list(dict.fromkeys(spec[col][0] + allowed)), c)
                    else:
                        spec[col] = (allowed, c)
    return spec


# ----------------------------------------------------------------- cleaning
def _norm(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def clean(df: pd.DataFrame, g: dict, problems: list) -> pd.DataFrame:
    idc, ep = g["id_column"], g["episode_column"]
    for col in (idc, ep):
        if col not in df.columns:
            raise SystemExit(f"Column '{col}' not found. Columns: {list(df.columns)}")
    df = df.copy()
    for col in df.columns:
        df[col] = df[col].map(_norm)
    df = df[df[idc] != ""].copy()
    dup = df[idc].str.casefold()[df[idc].str.casefold().duplicated()].unique().tolist()
    if dup:
        raise SystemExit(f"Duplicate entry IDs (case-insensitive): {dup}. Fix the sheet; the run was stopped so no count is produced.")
    # merge episode IDs that differ only in case/spacing
    key = df[ep].str.casefold()
    for k, grp in df[df[ep] != ""].groupby(key):
        forms = sorted(set(grp[ep]))
        if len(forms) > 1:
            problems.append({"level": "ERROR", "where": forms[0], "problem": f"Episode ID written in different forms {forms}; merged as '{forms[0]}'. Correct the sheet."})
            df.loc[grp.index, ep] = forms[0]
    # validate coded values
    spec = column_specs(g)
    for col in criterion_columns(g):
        if col not in df.columns:
            if g.get("frozen"):
                raise SystemExit(f"Frozen gates require column '{col}', which is missing from the data. A frozen run cannot proceed with missing criteria.")
            problems.append({"level": "ERROR", "where": "columns", "problem": f"Criterion column '{col}' is missing from the data (treated as Unknown; exploratory only)"})
            df[col] = ""
            continue
        allowed, c = spec[col]
        if allowed is None:
            bad = df[(df[col] != "") & pd.to_numeric(df[col], errors="coerce").isna()]
            for i in bad.index:
                problems.append({"level": "ERROR", "where": df.at[i, idc], "problem": f"'{col}' = '{df.at[i, col]}' is not a number (treated as Unknown)"})
                df.at[i, col] = ""
            continue
        canon = {a.casefold(): a for a in allowed}
        for i, v in df[col].items():
            if v == "":
                continue
            if v.casefold() in canon:
                if v != canon[v.casefold()]:
                    df.at[i, col] = canon[v.casefold()]
                continue
            problems.append({"level": "ERROR", "where": df.at[i, idc], "problem": f"'{col}' = '{v}' is not an allowed value {allowed} (treated as Unknown)"})
            df.at[i, col] = ""
    return df


# ----------------------------------------------------------------- evaluation
def eval_criterion(c: dict, row, g: dict) -> str:
    yes, no, unk = (set(x) for x in _vals(g))
    t = c.get("type", "yn")
    vals = [_norm(row.get(col)) for col in _cols(c)]
    if any(v == CONFLICT for v in vals):
        return HOLD
    if t == "yn":
        v = vals[0]
        return PASS if v in yes else FAIL if v in no else HOLD
    if t == "in":
        v = vals[0]
        if v == "" or v in unk:
            return HOLD
        return PASS if v in c["allowed"] else FAIL
    if t == "any_of":
        if any(v in yes for v in vals):
            return PASS
        return FAIL if all(v in no for v in vals) else HOLD
    if t == "min":
        v = pd.to_numeric(vals[0], errors="coerce")
        return HOLD if pd.isna(v) else PASS if v >= c["value"] else FAIL
    raise SystemExit(f"Unknown criterion type {t}")


def eval_stage(stage: dict, row, g: dict) -> dict:
    """Score all criteria on all paths. A path passes if all its GATE criteria pass.
    The best path is reported (Pass < Hold < Fail); diagnostics for every path are kept."""
    results, best = [], None
    for name, crits in zip(_path_names(stage), _paths(stage)):
        res = {c["id"]: eval_criterion(c, row, g) for c in crits}
        gates = [c for c in crits if c.get("gate", True)]
        fails = [c["id"] for c in gates if res[c["id"]] == FAIL]
        holds = [c["id"] for c in gates if res[c["id"]] == HOLD]
        status = FAIL if fails else HOLD if holds else PASS
        out = {"status": status, "path": name, "first_failed": fails[0] if fails else "", "first_held": holds[0] if holds else "",
               "failed": ";".join(fails), "missing": ";".join(holds), "scores": res,
               "flags": {c["id"]: res[c["id"]] for c in crits if not c.get("gate", True)}}
        results.append(out)
        rank = {PASS: 0, HOLD: 1, FAIL: 2}[status]
        if best is None or rank < best[0]:
            best = (rank, out)
    out = dict(best[1])
    if len(results) > 1:
        out["all_paths"] = " || ".join(f"{r['path']}: {r['status']}" + (f" fail[{r['failed']}]" if r["failed"] else "") + (f" held[{r['missing']}]" if r["missing"] else "") for r in results)
    return out


def _stage_cols(prefix, res, frame, multi):
    frame[f"{prefix}_status"] = [x["status"] if x else NA for x in res]
    frame[f"{prefix}_path"] = [x["path"] if x and x["status"] == PASS else "" for x in res]
    frame[f"{prefix}_first_failed"] = [x["first_failed"] if x else "" for x in res]
    frame[f"{prefix}_first_held"] = [x["first_held"] if x else "" for x in res]
    frame[f"{prefix}_failed"] = [x["failed"] if x else "" for x in res]
    frame[f"{prefix}_missing"] = [x["missing"] if x else "" for x in res]
    if multi:
        frame[f"{prefix}_all_paths"] = [x.get("all_paths", "") if x else "" for x in res]


def apply_gates(df: pd.DataFrame, g: dict):
    idc, ep = g["id_column"], g["episode_column"]
    problems: list = []
    df = clean(df, g, problems)
    entry_stages = [s for s in g["stages"] if s.get("unit", "episode") == "entry"]
    ep_stages = [s for s in g["stages"] if s.get("unit", "episode") == "episode"]

    # ---- entry-level stages (S0)
    df["_entry_state"] = PASS
    for st in entry_stages:
        res = [eval_stage(st, r, g) if r["_entry_state"] == PASS else None for _, r in df.iterrows()]
        _stage_cols(st["id"], res, df, "paths" in st)
        df["_entry_state"] = [s if s != NA else prev for s, prev in zip(df[f"{st['id']}_status"], df["_entry_state"])]
    no_ep = df[(df["_entry_state"] != FAIL) & (df[ep] == "")]
    for i in no_ep.index:
        problems.append({"level": "ERROR", "where": df.at[i, idc], "problem": "Entry is not excluded at Stage 0 but has no episode ID"})

    # consistency of episode-property criteria at S0 across entries of one episode
    entry_specific = {c["id"] for st in entry_stages for crits in _paths(st) for c in crits if c.get("entry_specific")}
    for st in entry_stages:
        for crits in _paths(st):
            for c in crits:
                if c["id"] in entry_specific:
                    continue
                for eid, grp in df[df[ep] != ""].groupby(ep):
                    vals = sorted(set(grp[c["column"]]) - {""})
                    if len(vals) > 1:
                        problems.append({"level": "ERROR", "where": eid, "problem": f"Entries of this episode disagree on {c['id']} ({c.get('label', '')}): {vals}"})

    # ---- episodes: built from all entries that are not excluded at S0; S0 status aggregated
    ep_cols = criterion_columns(g, "episode")
    rows = []
    for eid, grp in df[df[ep] != ""].groupby(ep, sort=False):
        states = set(grp["_entry_state"])
        s0 = PASS if PASS in states else HOLD if HOLD in states else FAIL
        usable = grp[grp["_entry_state"] != FAIL]
        r = {ep: eid, "S0_episode_status": s0, "n_entries": len(grp),
             "entries_used": ";".join(usable[idc]), "entries_excluded_S0": ";".join(grp.loc[grp["_entry_state"] == FAIL, idc]),
             "entries_held_S0": ";".join(grp.loc[grp["_entry_state"] == HOLD, idc])}
        if g.get("round_column") in grp:
            r["rounds"] = ";".join(sorted(set(grp[g["round_column"]]) - {""}))
        for c in ep_cols:
            vals = sorted(set(usable[c]) - {""}) if c in usable else []
            placeholders = {str(x) for x in g.get("values", {}).get("unknown", [])}
            dropped = sorted(set(grp.loc[grp["_entry_state"] == FAIL, c]) - {""} - placeholders) if c in grp else []
            # only a problem when the episode keeps other entries; an episode excluded as a whole simply sits outside
            if dropped and not vals and len(usable):
                problems.append({"level": "ERROR", "where": str(eid), "problem": f"'{c}' is coded only on entries excluded at Stage 0 ({r['entries_excluded_S0']}); move it to a retained entry"})
            if len(vals) > 1:
                problems.append({"level": "ERROR", "where": str(eid), "problem": f"Conflicting values for '{c}' across entries: {vals} (episode held on this criterion)"})
            r[c] = vals[0] if len(vals) == 1 else ("" if not vals else CONFLICT)
        for extra in g.get("carry_columns", []) + [x for x in [g.get("reason_column"), g.get("claimed_status_column")] if x]:
            if extra in grp:
                r[extra] = " | ".join(sorted(set(grp[extra]) - {""}))
        rows.append(r)
    eps = pd.DataFrame(rows)
    if not len(eps):
        eps = pd.DataFrame(columns=[ep, "S0_episode_status", "final_status"] + (["level"] if g.get("levels") else []))
        problems.append({"level": "ERROR", "where": "episodes", "problem": "No episodes could be formed from the data"})

    alive = pd.Series(True, index=eps.index) if len(eps) else pd.Series(dtype=bool)
    if len(eps):
        alive = eps["S0_episode_status"] == PASS
        for st in ep_stages:
            res = [eval_stage(st, r, g) if a else None for (_, r), a in zip(eps.iterrows(), alive)]
            _stage_cols(st["id"], res, eps, "paths" in st)
            for cid in dict.fromkeys(c["id"] for crits in _paths(st) for c in crits):
                eps[f"{st['id']}:{cid}"] = [x["scores"].get(cid, "") if x else "" for x in res]
            alive = alive & (eps[f"{st['id']}_status"] == PASS)

    levels = g.get("levels")

    def final(r):
        if r["S0_episode_status"] == FAIL:
            return "Outside analytic archive (S0)"
        if r["S0_episode_status"] == HOLD:
            return "Held at S0 (access/evidence)"
        for st in ep_stages:
            s = r[f"{st['id']}_status"]
            if s == FAIL:
                return f"Not passed {st['id']}"
            if s == HOLD:
                return f"Held at {st['id']} (evidence missing)"
        last = ep_stages[-1]
        return f"Passed {last['id']}" + (f" ({r[last['id'] + '_path']})" if "paths" in last else "")

    def level(r):
        if r["S0_episode_status"] == FAIL:
            return "Outside analytic archive"
        if r["S0_episode_status"] == HOLD:
            return "Held at S0 (not yet analysable)"
        lab = levels.get(entry_stages[-1]["id"] if entry_stages else "base", "Level 1")
        for st in ep_stages:
            s = r[f"{st['id']}_status"]
            if s == PASS:
                by_path = st.get("levels_by_path", {})
                lab = by_path.get(r[f"{st['id']}_path"], levels.get(st["id"], lab))
                continue
            if s == HOLD:
                lab += f" (pending evidence at {st['id']})"
            break
        return lab

    if len(eps):
        eps.insert(1, "final_status", eps.apply(final, axis=1))
        if levels:
            eps.insert(1, "level", eps.apply(level, axis=1))
    eps = eps.loc[:, ~eps.columns.duplicated()]

    # ---- written reasons (aggregated into one line per stage)
    rc = g.get("reason_column")
    if rc and len(eps) and rc in eps:
        miss = eps.loc[eps["final_status"].str.startswith(("Not passed", "Outside")) & (eps[rc].fillna("") == ""), ep].tolist()
        if miss:
            problems.append({"level": "WARN", "where": f"{len(miss)} episodes", "problem": "No written reason for not reaching the next level: " + ", ".join(map(str, miss))})

    # ---- hand-typed status vs gates
    cs = g.get("claimed_status_column")
    if cs and len(eps) and cs in eps:
        for _, r in eps.iterrows():
            c = (r.get(cs) or "").strip()
            if not c:
                continue
            if levels and c.lower().startswith("level"):
                got = r["level"].split(" (")[0]
                if not got.lower().startswith(c.split(" (")[0].lower()) and not c.lower().startswith(got.lower()):
                    problems.append({"level": "ERROR", "where": r[ep], "problem": f"Hand-typed level '{c}' disagrees with the gates: '{r['level']}'"})
            elif ("core" in c.lower()) != r["final_status"].startswith("Passed"):
                problems.append({"level": "ERROR", "where": r[ep], "problem": f"Hand-typed status '{c}' disagrees with the gates: '{r['final_status']}'"})

    df = df.drop(columns="_entry_state")
    return df, eps, pd.DataFrame(problems, columns=["level", "where", "problem"])


# ----------------------------------------------------------------- funnel
def criterion_labels(g):
    out = {}
    for st in g["stages"]:
        for crits in _paths(st):
            for c in crits:
                out[c["id"]] = c.get("label", c["id"])
    return out


def funnel(df, eps, g):
    rows = []
    rnd = g.get("round_column")
    lab = criterion_labels(g)
    nested = bool(g.get("levels"))

    def add(stage, kind, label, n, sub=None):
        r = {"stage": stage, "kind": kind, "count": label, "n": int(n)}
        if sub is not None and rnd and rnd in sub:
            r.update({f"n[{k or 'no round'}]": int(v) for k, v in sub[rnd].value_counts().items()})
        rows.append(r)

    add("entries", "main", "Source entries listed", len(df), df)
    for st in [s for s in g["stages"] if s.get("unit", "episode") == "entry"]:
        s = df[f"{st['id']}_status"]
        for fc, n in df.loc[s == FAIL, f"{st['id']}_first_failed"].value_counts().items():
            add(st["id"], "side", f"Entries excluded: {lab.get(fc, fc)}", n, df[(s == FAIL) & (df[f"{st['id']}_first_failed"] == fc)])
        for fc, n in df.loc[s == HOLD, f"{st['id']}_first_held"].value_counts().items():
            add(st["id"], "side", f"Entries held: {lab.get(fc, fc)}", n)
        add(st["id"], "main", f"Entries passing {st['id']} ({st['name']})", (s == PASS).sum(), df[s == PASS])
    if len(eps):
        add("episodes", "main", "Distinct episodes (all entries)", len(eps))
        for k, name in ((FAIL, "Episodes outside analytic archive"), (HOLD, "Episodes held at S0 (access/evidence)")):
            n = (eps["S0_episode_status"] == k).sum()
            if n:
                add("episodes", "side", name, n)
        add("episodes", "main", "Episodes in analytic archive", (eps["S0_episode_status"] == PASS).sum())
        verb = "Remain at lower level" if nested else "Excluded"
        for st in [s for s in g["stages"] if s.get("unit", "episode") == "episode"]:
            s = eps[f"{st['id']}_status"]
            for fc, n in eps.loc[s == FAIL, f"{st['id']}_first_failed"].value_counts().items():
                add(st["id"], "side", f"{verb}: {lab.get(fc, fc)}", n)
            for fc, n in eps.loc[s == HOLD, f"{st['id']}_first_held"].value_counts().items():
                add(st["id"], "side", f"Pending evidence: {lab.get(fc, fc)}", n)
            if "paths" in st:
                for pth, n in eps.loc[s == PASS, f"{st['id']}_path"].value_counts().items():
                    add(st["id"], "side", f"Passing via path: {pth}", n)
            add(st["id"], "main", f"Episodes passing {st['id']} ({st['name']})", (s == PASS).sum())
        if "level" in eps:
            base = eps["level"].str.replace(r" \(pending.*\)$", "", regex=True)
            for lv, n in base.value_counts().sort_index().items():
                add("levels", "level", lv, n)
            for lv, n in eps.loc[eps["level"].str.contains("pending"), "level"].value_counts().sort_index().items():
                add("levels", "level", f"  of which {lv}", n)
    return pd.DataFrame(rows)


def flow_svg(fun: pd.DataFrame, g: dict) -> str:
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;")
    W, bw, bh, gapy = 1080, 400, 58, 50
    y, parts, pending, first = 30, [], [], True
    for _, m in fun[fun["kind"].isin(["main", "side"])].iterrows():
        if m["kind"] == "side":
            pending.append(m)
            continue
        if not first:
            top = y - gapy
            if pending:
                sy, h = top + 4, 18 * len(pending) + 14
                parts.append(f'<rect x="480" y="{sy}" width="580" height="{h}" rx="4" fill="#FFF7F2" stroke="#B4441B"/>')
                for i, s in enumerate(pending):
                    parts.append(f'<text x="490" y="{sy + 20 + 18 * i}" font-size="12" fill="#5a2a14">{esc(s["count"])} (n = {s["n"]})</text>')
                parts.append(f'<line x1="{40 + bw / 2}" y1="{sy + h / 2}" x2="480" y2="{sy + h / 2}" stroke="#B4441B"/>')
                y = max(y, sy + h + 8)
            parts.append(f'<line x1="{40 + bw / 2}" y1="{top}" x2="{40 + bw / 2}" y2="{y - 2}" stroke="#132A44" marker-end="url(#a)"/>')
        pending, first = [], False
        parts.append(f'<rect x="40" y="{y}" width="{bw}" height="{bh}" rx="6" fill="#EFF4F7" stroke="#132A44"/>')
        parts.append(f'<text x="{40 + bw / 2}" y="{y + 25}" text-anchor="middle" font-size="14" font-weight="bold" fill="#132A44">{esc(m["count"])}</text>')
        parts.append(f'<text x="{40 + bw / 2}" y="{y + 45}" text-anchor="middle" font-size="15" fill="#132A44">n = {m["n"]}</text>')
        y += bh + gapy
    lv = fun[fun["kind"] == "level"]
    if len(lv):
        parts.append(f'<text x="40" y="{y}" font-size="14" font-weight="bold" fill="#132A44">Analysis level assigned (no episode removed)</text>')
        y += 10
        for _, r in lv.iterrows():
            y += 22
            parts.append(f'<rect x="40" y="{y - 16}" width="620" height="20" fill="#E2F0D9"/><text x="48" y="{y}" font-size="13" fill="#1F4E2C">{esc(r["count"])}: n = {r["n"]}</text>')
        y += 30
    head = f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{y + 10}" font-family="Arial"><defs><marker id="a" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="#132A44"/></marker></defs>'
    foot = f'<text x="40" y="{y}" font-size="11" fill="#667281">Gates {esc(g["protocol_version"])}{" — DRAFT, not frozen" if not g.get("frozen") else " — frozen " + esc(g.get("frozen_on", ""))}</text>'
    return head + "".join(parts) + foot + "</svg>"


# ----------------------------------------------------------------- IO
def load_data_sheet(path, sheet=None, header_row=0):
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(p, sheet_name=sheet or 0, header=header_row, dtype=str, keep_default_na=False, na_values=[""])
    return pd.read_csv(p, dtype=str, keep_default_na=False, na_values=[""])


def append_ledger(out_dir, entry):
    p = Path(out_dir) / "RUN_LEDGER.jsonl"
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")


# ----------------------------------------------------------------- commands
def cmd_run(a):
    g = load_gates(a.gates)
    df = load_data_sheet(a.data, a.sheet, a.header_row)
    with RunLog(a.out, "screen", "run", vars(a), None) as run:
        run.add_input(a.data)
        run.add_input(a.gates, "gates")
        h, citable = check_frozen(g, a.gates, run=run)
        entries, eps, probs = apply_gates(df, g)
        n_err = int((probs["level"] == "ERROR").sum())
        citable = citable and n_err == 0
        run.note("gates", {"version": g["protocol_version"], "frozen": bool(g.get("frozen")), "sha256": h})
        run.note("citable", citable)
        fun = funnel(entries, eps, g)
        crit = pd.DataFrame([{"id": k, "label": v} for k, v in criterion_labels(g).items()])
        out = run.path("screening_result.xlsx")
        with pd.ExcelWriter(out, engine="openpyxl") as xw:
            pd.DataFrame({"status": ["CITABLE" if citable else "NOT CITABLE — draft gates and/or ERRORs in Checks"],
                          "gates": [f"{g['protocol_version']} sha256 {h[:16]}"], "data_sha256": [sha256_file(a.data)[:16]]}).to_excel(xw, sheet_name="Status", index=False)
            fun.to_excel(xw, sheet_name="Funnel", index=False)
            eps.to_excel(xw, sheet_name="Episodes", index=False)
            entries.to_excel(xw, sheet_name="Entries", index=False)
            probs.to_excel(xw, sheet_name="Checks", index=False)
            crit.to_excel(xw, sheet_name="Criteria", index=False)
        run.add_output(out)
        svg = run.path("flow_diagram.svg")
        svg.write_text(flow_svg(fun, g), encoding="utf-8")
        run.add_output(svg)
        md = run.path("funnel.md")
        md.write_text(fun.to_markdown(index=False) + f"\n\nChecks: {n_err} errors, {int((probs['level'] == 'WARN').sum())} warnings. Citable: {citable}\n", encoding="utf-8")
        run.add_output(md)
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            main = fun[fun["kind"].isin(["main", "level"])]
            fig, ax = plt.subplots(figsize=(8, 0.42 * len(main) + 0.8))
            cols = ["#A4772B" if "Level 3" in str(c) else "#2A7A6D" if "Level 2" in str(c) else "#5A7696" if "Level 1" in str(c) else "#14233A" for c in main["count"]]
            ax.barh(range(len(main))[::-1], main["n"], color=cols)
            ax.set_yticks(range(len(main))[::-1], [str(c).strip()[:60] for c in main["count"]], fontsize=8)
            for i, v in zip(range(len(main))[::-1], main["n"]):
                ax.text(v, i, f" {v}", va="center", fontsize=8)
            ax.spines[["top", "right"]].set_visible(False)
            ax.set_xlabel("n")
            fig.tight_layout()
        except Exception:
            fig = None
        from ..pdfreport import add_pdf
        lv = eps["level"].value_counts().rename_axis("level").reset_index(name="episodes") if "level" in eps.columns else None
        add_pdf(run, "screening_report.pdf", "Screening result", [
            ("kv", [("Status", "CITABLE" if citable else "NOT CITABLE (draft gates and/or errors)"),
                    ("Protocol", f"{g['protocol_version']} · sha256 {h[:16]}"), ("Data sha256", sha256_file(a.data)[:16]),
                    ("Errors / warnings", f"{n_err} / {int((probs['level'] == 'WARN').sum())}")]),
            ("h", "Screening flow"), *([("image", fig)] if fig is not None else []), ("table", fun[["stage", "count", "n"]]),
            ("h", "Episodes by level"), ("table", lv),
            ("h", "Checks"), ("table", probs if len(probs) else None)], subtitle=Path(a.data).name)
        if fig is not None:
            plt.close(fig)
        append_ledger(a.out, {"time": _dt.datetime.now().isoformat(timespec="seconds"), "run_folder": str(run.dir), "data": str(Path(a.data).resolve()),
                              "data_sha256": sha256_file(a.data), "gates_version": g["protocol_version"], "gates_sha256": h, "citable": citable,
                              "counts": {r["count"].strip(): r["n"] for _, r in fun.iterrows() if r["kind"] in ("main", "level")}})
        print(fun[["stage", "count", "n"]].to_string(index=False))
        print(f"\n{n_err} errors, {int((probs['level'] == 'WARN').sum())} warnings — see Checks sheet. Citable: {citable}")
        run.note("errors", n_err)
        run.note("result_status", "completed_with_errors" if n_err else "completed_clean")
    if n_err:
        raise SystemExit(2)


def cmd_compare(a):
    ga, gb = load_gates(a.gates), load_gates(a.gates_b)
    df = load_data_sheet(a.data, a.sheet, a.header_row)
    with RunLog(a.out, "screen", "compare", vars(a), None) as run:
        run.add_input(a.data)
        run.add_input(a.gates, "gates_a")
        run.add_input(a.gates_b, "gates_b")
        _, ea, _ = apply_gates(df, ga)
        _, eb, _ = apply_gates(df, gb)
        ep = ga["episode_column"]
        col = lambda e: "level" if "level" in e else "final_status"
        A = ea[[ep, col(ea)]].rename(columns={col(ea): f"[{ga['protocol_version']}]"})
        B = eb[[ep, col(eb)]].rename(columns={col(eb): f"[{gb['protocol_version']}]"})
        m = pd.merge(A, B, on=ep, how="outer").fillna("not an episode")
        c1, c2 = [c for c in m.columns if c != ep]
        m["changed"] = m[c1] != m[c2]
        out = run.path("gate_comparison.xlsx")
        m.sort_values("changed", ascending=False).to_excel(out, index=False)
        run.add_output(out)
        print(m[m.changed].to_string(index=False) if m.changed.any() else "No episode changes status.")


def cmd_template(a):
    g = load_gates(a.gates)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = [g["id_column"], g["episode_column"]] + [c for c in [g.get("round_column"), "register_source", "source_author", "description_pseudonymised"] if c]
    cols += criterion_columns(g)
    cols += [c for c in [g.get("reason_column"), g.get("claimed_status_column")] if c] + g.get("carry_columns", [])
    extras = {e["column"]: e for e in g.get("extra_columns", [])}
    cols += list(extras)
    cols = list(dict.fromkeys(cols))
    spec = column_specs(g)
    wb = Workbook()
    ws = wb.active
    ws.title = "Screening"
    for j, col in enumerate(cols, 1):
        cell = ws.cell(1, j, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="132A44")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = 18
        if col in spec:
            allowed, c = spec[col]
            d = ws.cell(2, j, f"{c['id']} · {c.get('label', c['id'])}" + ("  [flag, not a gate]" if not c.get("gate", True) else ""))
            d.alignment = Alignment(wrap_text=True, vertical="top")
            d.font = Font(italic=True, size=9, color="667281")
            if allowed:
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{cell.column_letter}3:{cell.column_letter}1000")
        elif col in extras:
            e = extras[col]
            d = ws.cell(2, j, e.get("group", "") + (" · " + e["hint"] if e.get("hint") else ""))
            d.alignment = Alignment(wrap_text=True, vertical="top")
            d.font = Font(italic=True, size=9, color="1F4E2C")
            cell.fill = PatternFill("solid", fgColor="1F4E2C")
            if e.get("values"):
                dv = DataValidation(type="list", formula1='"' + ",".join(e["values"]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{cell.column_letter}3:{cell.column_letter}1000")
    ws.row_dimensions[2].height = 60
    ws.freeze_panes = "C3"
    d = wb.create_sheet("Definitions")
    d.append(["Stage", "Unit", "Path", "Criterion ID", "Gate?", "Column", "Label", "Definition / evidence required"])
    for st in g["stages"]:
        for pname, crits in zip(_path_names(st), _paths(st)):
            for c in crits:
                d.append([st["id"], st.get("unit", "episode"), pname, c["id"], "yes" if c.get("gate", True) else "flag", ", ".join(_cols(c)), c.get("label", ""), c.get("definition", "")])
    wb.save(a.output)
    print(f"Template written: {a.output}  (row 2 = criterion labels, ignored because it has no entry ID; enter data from row 3)")


def cmd_freeze(a):
    g = load_gates(a.gates)
    if g.get("frozen"):
        raise SystemExit(f"Already frozen on {g.get('frozen_on')} by {g.get('frozen_by')}.")
    if "draft" in g["protocol_version"].lower():
        raise SystemExit("Rename protocol_version without 'DRAFT' before freezing (e.g. 'v3').")
    if any(r["protocol_version"] == g["protocol_version"] for r in read_register(a.gates)):
        raise SystemExit(f"{g['protocol_version']} was frozen before. Use a new protocol_version.")
    g["frozen"] = True
    g["frozen_on"] = _dt.datetime.now().isoformat(timespec="seconds")
    g["frozen_by"] = a.by
    g["sha256"] = gates_hash(g)
    with open(a.gates, "w", encoding="utf-8") as f:
        json.dump(g, f, indent=2, ensure_ascii=False)
    with open(register_path(a.gates), "a", encoding="utf-8") as f:
        f.write(json.dumps({"protocol_version": g["protocol_version"], "sha256": g["sha256"], "frozen_on": g["frozen_on"], "frozen_by": a.by, "file": Path(a.gates).name}, ensure_ascii=False) + "\n")
    print(f"Frozen {g['protocol_version']} at {g['frozen_on']} by {a.by}. sha256 {g['sha256']}")
    print(f"Registered in {register_path(a.gates)}.")
    print("For an external timestamp, upload the gates file and FREEZE_REGISTER.jsonl to OSF/Zenodo now.")
    print(f'Record: "USER DECISION — {g["frozen_on"][:10]}: screening gates {g["protocol_version"]} frozen (sha256 {g["sha256"][:12]})."')


def _agreement_stats(x, y, scale):
    cats = list(dict.fromkeys(list(scale) + sorted(set(x) | set(y))))
    n = len(x)
    if n == 0:
        return np.nan, np.nan, np.nan
    po = float(np.mean([a == b for a, b in zip(x, y)]))
    px = {c: np.mean([a == c for a in x]) for c in cats}
    py = {c: np.mean([b == c for b in y]) for c in cats}
    pe = sum(px[c] * py[c] for c in cats)
    kappa = (po - pe) / (1 - pe) if pe < 1 else np.nan
    q = len(cats)
    pi = {c: (px[c] + py[c]) / 2 for c in cats}
    pe_ac1 = sum(p * (1 - p) for p in pi.values()) / (q - 1) if q > 1 else np.nan
    ac1 = (po - pe_ac1) / (1 - pe_ac1) if q > 1 and pe_ac1 < 1 else np.nan
    return po, kappa, ac1


def cmd_agreement(a):
    g = load_gates(a.gates)
    A = load_data_sheet(a.data, a.sheet, a.header_row)
    B = load_data_sheet(a.coder, a.sheet, a.header_row)
    with RunLog(a.out, "screen", "agreement", vars(a), None) as run:
        run.add_input(a.data, "author")
        run.add_input(a.coder, "coder")
        run.add_input(a.gates, "gates")
        h, _ = check_frozen(g, a.gates, run=run)
        run.note("gates", {"version": g["protocol_version"], "sha256": h})
        _, ea, pa = apply_gates(A, g)
        _, eb, pb = apply_gates(B, g)
        for who, p in (("author", pa), ("coder", pb)):
            if (p["level"] == "ERROR").any():
                run.warn(f"{who} file has {(p['level'] == 'ERROR').sum()} ERRORs (conflicts are compared as 'CONFLICT').")
        key = g["episode_column"]
        ea, eb = ea.set_index(key), eb.set_index(key)
        ea.index, eb.index = ea.index.str.casefold(), eb.index.str.casefold()
        ids = ea.index.intersection(eb.index)
        only_a, only_b = sorted(set(ea.index) - set(ids)), sorted(set(eb.index) - set(ids))
        if only_a or only_b:
            run.warn(f"Units only in author file: {only_a}; only in coder file: {only_b}")
        spec = column_specs(g)
        rows, dis = [], []
        cols = [c for c in criterion_columns(g, "episode") if c in ea.columns and c in eb.columns]
        extra = [c for c in ("level", "final_status") if c in ea.columns and c in eb.columns]
        for c in cols + extra:
            x = [(v or "Unknown") for v in ea.loc[ids, c].fillna("")]
            y = [(v or "Unknown") for v in eb.loc[ids, c].fillna("")]
            scale = (spec.get(c, (None,))[0] or [])
            po, k, ac1 = _agreement_stats(x, y, scale)
            rows.append({"column": c, "units": len(ids), "percent_agreement": round(100 * po, 1),
                         "cohen_kappa": None if k != k else round(k, 3), "gwet_ac1": None if ac1 != ac1 else round(ac1, 3),
                         "scale": ", ".join(scale) if scale else "observed"})
            dis += [{"unit": i, "column": c, "author": xa, "coder": yb, "resolution_note": ""} for i, xa, yb in zip(ids, x, y) if xa != yb]
        out = run.path("agreement.xlsx")
        with pd.ExcelWriter(out, engine="openpyxl") as xw:
            pd.DataFrame(rows).to_excel(xw, sheet_name="Agreement", index=False)
            pd.DataFrame(dis, columns=["unit", "column", "author", "coder", "resolution_note"]).to_excel(xw, sheet_name="Disagreements", index=False)
            pd.DataFrame({"only_in_author": pd.Series(only_a, dtype=str), "only_in_coder": pd.Series(only_b, dtype=str)}).to_excel(xw, sheet_name="Unmatched", index=False)
        run.add_output(out)
        print(pd.DataFrame(rows).to_string(index=False))
        if len(ids) < 10:
            run.warn(f"Only {len(ids)} units: report percent agreement and the disagreement list; kappa is unstable at this size.")


class Screen(Plugin):
    name = "screen"
    help = "Case screening with frozen gates: levels, funnel, flow diagram, integrity checks, gate comparison, coder agreement"

    def register(self, sub):
        def common(p, data=True):
            if data:
                p.add_argument("--data", required=True, help="Screening sheet (xlsx or csv), one row per source entry")
                p.add_argument("--sheet", default=None, help="Excel sheet name (default: first)")
                p.add_argument("--header-row", type=int, default=0, help="0-based header row (default 0)")
            p.add_argument("--gates", required=True, help="Gates JSON")
            p.add_argument("--out", default="clockbind_runs")
            return p

        p = common(sub.add_parser("run", help="Apply gates to every row; write levels, funnel, flow diagram, checks, ledger"))
        p.set_defaults(func=cmd_run)
        p = common(sub.add_parser("compare", help="Show which episodes change status/level between two gate versions"))
        p.add_argument("--gates-b", required=True)
        p.set_defaults(func=cmd_compare)
        p = common(sub.add_parser("agreement", help="Author vs independent coder: % agreement, Cohen's kappa, Gwet's AC1 (full scale), disagreements"))
        p.add_argument("--coder", required=True)
        p.set_defaults(func=cmd_agreement)
        p = common(sub.add_parser("template", help="Blank screening workbook with dropdowns from the gates"), data=False)
        p.add_argument("--output", default="screening_template.xlsx")
        p.set_defaults(func=cmd_template)
        p = common(sub.add_parser("freeze", help="Lock the gates once per protocol_version (append-only register)"), data=False)
        p.add_argument("--by", required=True)
        p.set_defaults(func=cmd_freeze)


PLUGIN = Screen()
