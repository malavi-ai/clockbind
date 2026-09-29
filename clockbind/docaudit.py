"""Document audit: personal data and outdated wording in Word, PDF, Excel, CSV and text files, all on this computer.

Personal-data findings report only the kind and the location, never the value. Wording findings report the
matched rule phrase and the current wording. Nothing leaves the computer.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .privacy import IBAN, _hu_tax_ok, _iban_ok, _luhn_ok, _tckn_ok

TEXT_TYPES = {".docx", ".docm", ".pdf", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}
ARCHIVE_TYPES = {".zip"}
MAX_ARCHIVE_FILES = 5000
MAX_MEMBER_BYTES = 100 * 1024 * 1024
MAX_TOTAL_ARCHIVE_BYTES = 500 * 1024 * 1024

EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"(?<![\w.])(?:\+|00)\d[\d ()./-]{7,}\d(?![\w])")
ELEVEN = re.compile(r"(?<!\d)\d{11}(?!\d)")
TEN = re.compile(r"(?<!\d)\d{10}(?!\d)")
CARD = re.compile(r"(?<!\d)(?:\d{4}[ -]){3}\d{1,7}(?!\d)|(?<!\d)\d{13,19}(?!\d)")
CUR = r"(?:HUF|Ft|EUR|€|USD|US\$|\$|TRY|TL|₺|RON|lei|GBP|£)"
AMOUNT = re.compile(rf"(?<![\w]){CUR}\s?\d[\d\s.,']*\d|(?<![\w.,])\d[\d\s.,']*\d\s?{CUR}(?![\w])", re.I)
INVOICE = re.compile(r"\b(?:invoice|inv|proforma|pro-forma|számla|szamla|fatura|factura|order|PO)\b\.?\s*(?:no|nr|number|sz|num|#)?\.?\s*[:#]?\s*"
                     r"(?=[A-Z0-9/\-]*\d)[A-Z0-9][A-Z0-9/\-]{3,}", re.I)
NAME_LABEL = re.compile(
    r"\b(?:customer|client|contact|buyer|owner|name|full\s+name|first\s+name|last\s+name|surname|"
    r"müşteri|musteri|ad(?:ı)?|soyad(?:ı)?|isim|név|ugyfel|ügyfél|nome|cognome|cliente|"
    r"نام|نام\s+خانوادگی|مشتری)\b\s*(?:[:=]|\s-\s)\s*"
    r"(?=[^\n]{2,80}$)(?:[A-ZÀ-ÖØ-öø-ÿĀ-žĞİŞÇÖÜ][\wÀ-ÖØ-öø-ÿĀ-žĞİŞÇÖÜ\'’.-]+(?:\s+|$)){1,4}",
    re.I,
)


NOISE = re.compile(r"https?://\S+|\bdoi:\s*\S+|\b10\.\d{4,9}/\S+|\bISBN[\s:-]*[\dXx-]{10,17}", re.I)


@dataclass
class Unit:
    file: str
    location: str
    text: str


# ------------------------------------------------------------------ reading
def _units_docx(data: bytes, name: str):
    import docx
    d = docx.Document(io.BytesIO(data))
    for i, p in enumerate(d.paragraphs, 1):
        if p.text.strip():
            yield Unit(name, f"paragraph {i}", p.text)
    for ti, t in enumerate(d.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                if cell.text.strip():
                    yield Unit(name, f"table {ti} row {ri} col {ci}", cell.text)
    for si, s in enumerate(d.sections, 1):
        for part, lab in ((s.header, "header"), (s.footer, "footer")):
            txt = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            if txt:
                yield Unit(name, f"{lab} (section {si})", txt)


def _units_pdf(data: bytes, name: str):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF reading needs the pypdf package:  pip install pypdf") from None
    r = PdfReader(io.BytesIO(data))
    for i, pg in enumerate(r.pages, 1):
        txt = pg.extract_text() or ""
        if txt.strip():
            yield Unit(name, f"page {i}", txt)


def _units_xlsx(data: bytes, name: str):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is not None and str(v).strip() and not hasattr(v, "year"):
                        yield Unit(name, f"{ws.title}!{c.coordinate}", str(v))
    finally:
        wb.close()


def _units_text(data: bytes, name: str):
    txt = data.decode("utf-8", errors="replace")
    for i, line in enumerate(txt.splitlines(), 1):
        if line.strip():
            yield Unit(name, f"line {i}", line)


READERS = {".docx": _units_docx, ".docm": _units_docx, ".pdf": _units_pdf, ".xlsx": _units_xlsx, ".xlsm": _units_xlsx,
           ".csv": _units_text, ".txt": _units_text, ".md": _units_text}


def _iter_zip_bytes(data: bytes, archive_name: str, depth: int = 0):
    """Yield supported members from a zip without extracting them. Nested zips are supported to depth 3.

    Limits are deliberately conservative so a malformed/hostile archive cannot expand without bound.
    """
    if depth > 3:
        yield f"{archive_name} › [nested archive depth exceeded]", ".skip", None
        return
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = [i for i in z.infolist() if not i.is_dir() and "__MACOSX" not in i.filename and not Path(i.filename).name.startswith(("~$", "."))]
        if len(infos) > MAX_ARCHIVE_FILES:
            raise RuntimeError(f"archive contains more than {MAX_ARCHIVE_FILES} files")
        for info in infos:
            if info.file_size > MAX_MEMBER_BYTES:
                yield f"{archive_name} › {info.filename}", ".skip", None
                continue
            total += info.file_size
            if total > MAX_TOTAL_ARCHIVE_BYTES:
                raise RuntimeError("archive expands beyond the local audit safety limit")
            n = info.filename
            suf = Path(n).suffix.lower()
            member = z.read(info)
            shown = f"{archive_name} › {n}"
            if suf in ARCHIVE_TYPES:
                yield from _iter_zip_bytes(member, shown, depth + 1)
            else:
                yield shown, suf, (member if suf in TEXT_TYPES else None)


def iter_files(path: str):
    """(display name, suffix, bytes) for a file, every supported file in a folder (recursively) or inside a zip."""
    p = Path(path).expanduser()
    if p.is_dir():
        for f in sorted(p.rglob("*")):
            if f.is_file() and not f.name.startswith(("~$", ".")):
                if f.suffix.lower() == ".zip":
                    yield from _iter_zip_bytes(f.read_bytes(), str(f.relative_to(p)))
                else:
                    yield from _one(f, str(f.relative_to(p)))
    elif p.suffix.lower() == ".zip":
        yield from _iter_zip_bytes(p.read_bytes(), p.name)
    else:
        yield from _one(p, p.name)

def _one(f: Path, name: str):
    suf = f.suffix.lower()
    yield name, suf, (f.read_bytes() if suf in TEXT_TYPES else None)


# ------------------------------------------------------------------ checks
def load_terms(path: str | None) -> dict:
    if path is None:
        return {"name": "(none)", "rules": [], "allow_amounts": []}
    g = json.loads(Path(path).read_text(encoding="utf-8"))
    for r in g.get("rules", []):
        r["_re"] = re.compile(r["pattern"], re.I if "i" in r.get("flags", "i") else 0)
    g["_allow"] = [re.compile(a, re.I) for a in g.get("allow_amounts", [])]
    g["_wording_exempt"] = [re.compile(a, re.I) for a in g.get("wording_exempt_locations", [])]
    return g


def load_names(path: str | None) -> list:
    if not path:
        return []
    names = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        n = line.strip().strip(",;")
        if len(n) >= 3 and not n.startswith("#"):
            names.append(re.compile(r"(?<!\w)" + re.escape(n) + r"(?!\w)", re.I))
    return names


def personal_kinds(text: str, names: list, allow: list) -> dict:
    """{kind: count} of personal or confidential data in a text unit. Values are not returned."""
    k: dict[str, int] = {}

    def add(kind, n=1):
        if n:
            k[kind] = k.get(kind, 0) + n

    add("e-mail address", len(EMAIL.findall(text)))
    text = NOISE.sub(" ", text)   # links, DOIs and ISBNs are not personal data
    add("phone number", sum(1 for m in PHONE.findall(text) if 9 <= len(re.sub(r"\D", "", m)) <= 15))
    add("bank account (IBAN)", sum(1 for m in IBAN.finditer(text.upper()) if _iban_ok(m.group(0))))
    add("Turkish ID number (TCKN)", sum(1 for m in ELEVEN.findall(text) if _tckn_ok(m)))
    add("Hungarian tax ID", sum(1 for m in TEN.findall(text) if _hu_tax_ok(m)))
    add("payment card number", sum(1 for m in CARD.findall(text) if _luhn_ok(re.sub(r"\D", "", m))))
    amount_text = text
    for al in allow:
        amount_text = al.sub(" ", amount_text)
    amounts = [m.group(0) for m in AMOUNT.finditer(amount_text)]
    add("exact amount", len(amounts))
    add("invoice or order number", len(INVOICE.findall(text)))
    add("listed name", sum(len(n.findall(text)) for n in names))
    # Conservative fallback for explicitly labelled name fields in prose. The matched value is never returned.
    add("person name (labelled field)", len(NAME_LABEL.findall(text)))
    return k


def audit(path: str, terms: dict, names: list) -> tuple[list, list, list]:
    """(findings, per-file summary, files skipped). Findings never contain personal values."""
    findings, summary, skipped = [], [], []
    for name, suf, data in iter_files(path):
        if data is None:
            note = "Google Docs shortcut: download it as .docx first" if suf in (".gdoc", ".gsheet") else "file type not read"
            skipped.append({"file": name, "reason": note})
            continue
        pd_n = term_n = units = 0
        try:
            for u in READERS[suf](data, name):
                units += 1
                for kind, n in personal_kinds(u.text, names, terms.get("_allow", [])).items():
                    findings.append({"file": name, "location": u.location, "check": "personal or confidential data",
                                     "finding": kind, "count": n, "severity": "HIGH" if kind in ("listed name", "e-mail address", "phone number",
                                     "bank account (IBAN)", "Turkish ID number (TCKN)", "Hungarian tax ID", "payment card number", "person name (labelled field)") else "MEDIUM",
                                     "current wording / action": "remove or replace with a pseudonymised code", "matched": ""})
                    pd_n += n
                wording_exempt = any(x.search(u.location) for x in terms.get("_wording_exempt", []))
                if not wording_exempt:
                    for r in terms.get("rules", []):
                        for m in r["_re"].finditer(u.text):
                            findings.append({"file": name, "location": u.location, "check": "outdated wording", "finding": r["issue"],
                                             "count": 1, "severity": r.get("severity", "MEDIUM"), "current wording / action": r.get("current", ""),
                                             "matched": m.group(0)[:80]})
                            term_n += 1
        except Exception as e:
            skipped.append({"file": name, "reason": f"could not be read ({type(e).__name__}: {e})"})
            continue
        if units == 0:
            skipped.append({"file": name, "reason": "no extractable text found (possibly scanned/image-only or empty)"})
            continue
        summary.append({"file": name, "text units read": units, "personal-data findings": pd_n, "outdated-wording findings": term_n,
                        "status": "clean" if not (pd_n or term_n) else ("personal data" if pd_n else "wording to update")})
    return findings, summary, skipped
