"""Independent-coder tools: blind form carries IDs only; agreement statistics and level comparison."""
from pathlib import Path
import argparse

import numpy as np
import pandas as pd
import pytest

openpyxl = pytest.importorskip("openpyxl")
from tests.test_fullreport import _workbook  # noqa: E402


def test_weighted_kappa_known_values():
    from clockbind.plugins.screen import weighted_kappa
    assert weighted_kappa([(0, 0), (1, 1), (2, 2), (1, 1)], 3) == pytest.approx(1.0)
    # hand-computed: o = [[.5,0],[.25,.25]] -> po_w=.75, pe_w=.5 -> kappa .5 (k=2, linear = unweighted)
    assert weighted_kappa([(0, 0), (0, 0), (1, 0), (1, 1)], 2) == pytest.approx(0.5)


def test_coder_form_is_blind_and_agreement_runs(tmp_path):
    from clockbind.plugins import screen
    from clockbind.plugins.report import DEFAULT_GATES
    author = tmp_path / "author.xlsx"
    _workbook(author)
    form = tmp_path / "form.xlsx"
    screen.cmd_template(argparse.Namespace(gates=str(DEFAULT_GATES), output=str(form), episodes_from=str(author),
                                           from_sheet="05_Level_Assessment", from_header_row=2, out=str(tmp_path / "runs")))
    from openpyxl import load_workbook
    ws = load_workbook(form)["Screening"]
    vals = [c.value for row in ws.iter_rows(min_row=3) for c in row if c.value not in (None, "")]
    assert sorted(vals) == ["ZZ1", "ZZ2"]                       # IDs only, nothing else copied
    hdr = [c.value for c in ws[1]]
    assert "Buyer stage" in hdr and "L3e-S source type" in hdr
    # coder copy = author with one gate changed
    coder = tmp_path / "coder.xlsx"
    wb = load_workbook(author)
    sh = wb["05_Level_Assessment"]
    import json
    g = json.loads(Path(DEFAULT_GATES).read_text())
    l3c = next(c["column"] for s in g["stages"] for c in s["criteria"] if c["id"] == "L3c")
    col = [c.value for c in sh[3]].index(l3c) + 1
    sh.cell(5, col, "No")                                          # ZZ2: coder says L3c = No
    wb.save(coder)
    a = argparse.Namespace(gates=str(DEFAULT_GATES), data=str(author), coder=str(coder), sheet="05_Level_Assessment",
                           header_row=2, out=str(tmp_path / "runs"), bootstrap=200, seed=1)
    screen.cmd_agreement(a)
    out = sorted((tmp_path / "runs").rglob("agreement.xlsx"))[-1]
    summ = pd.read_excel(out, sheet_name="Summary").set_index("measure")["value"]
    assert int(summ["Level 3 — author"]) == 1 and int(summ["Level 3 — coder"]) == 0
    assert int(summ["Units whose computed level differs"]) == 1
    ag = pd.read_excel(out, sheet_name="Agreement").set_index("column")
    assert ag.loc[l3c, "percent_agreement"] == 50.0
