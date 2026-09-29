import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pandas as pd


def _write_docx(path: Path, text: str):
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    d.save(path)


def test_audit_plugin_is_auto_discovered():
    from clockbind.core.plugin import discover
    names = {p.name for p in discover()}
    assert "audit" in names


def test_document_audit_cli_safe_summary_and_local_detail(tmp_path):
    root = Path(__file__).resolve().parents[1]
    docs = tmp_path / "docs"; docs.mkdir()
    sensitive = docs / "Jane Example invoice.docx"
    _write_docx(sensitive, "Name: Jane Example. Contact jane@example.com. Invoice no INV-2026-0042. A = shows.")
    summary = tmp_path / "summary.json"
    out = tmp_path / "runs"
    r = subprocess.run([
        sys.executable, "-m", "clockbind", "audit", "docs", "--data", str(docs),
        "--summary-json", str(summary), "--out", str(out)
    ], cwd=root, capture_output=True, text=True, env={**__import__('os').environ, "PYTHONPATH": str(root)})
    assert r.returncode == 0, r.stderr
    # Console is safe to paste into a chat: no file names and no matched personal values.
    assert "Jane Example" not in r.stdout and "jane@example.com" not in r.stdout
    assert sensitive.name not in r.stdout
    s = json.loads(summary.read_text(encoding="utf-8"))
    assert s == {"files_read": 1, "files_with_personal_or_confidential_data": 1,
                 "files_with_outdated_wording": 1, "files_not_read": 0}
    run_dir = next(out.iterdir())
    x = pd.ExcelFile(run_dir / "document_audit.xlsx")
    f = pd.read_excel(x, "Findings")
    assert {"e-mail address", "invoice or order number", "person name (labelled field)"} <= set(f["finding"])
    assert not f.loc[f["check"].eq("personal or confidential data"), "matched"].fillna("").astype(str).str.len().any()


def test_document_audit_cli_list_files_is_explicit_opt_in(tmp_path):
    root = Path(__file__).resolve().parents[1]
    f = tmp_path / "Local Person.docx"
    _write_docx(f, "Name: Local Person")
    r = subprocess.run([
        sys.executable, "-m", "clockbind", "audit", "docs", "--data", str(f), "--list-files",
        "--out", str(tmp_path / "runs")
    ], cwd=root, capture_output=True, text=True, env={**__import__('os').environ, "PYTHONPATH": str(root)})
    assert r.returncode == 0
    assert f.name in r.stdout


def test_document_audit_nested_zip_and_scanned_pdf_notice(tmp_path):
    from clockbind.docaudit import audit, load_terms
    from clockbind.resources import resource
    inner = tmp_path / "inner.zip"
    with zipfile.ZipFile(inner, "w") as z:
        z.writestr("note.txt", "Contact someone@example.org. A = shows.")
    outer = tmp_path / "outer.zip"
    with zipfile.ZipFile(outer, "w") as z:
        z.write(inner, "nested/inner.zip")
        z.writestr("shortcut.gdoc", "{}")
    findings, summary, skipped = audit(str(outer), load_terms(str(resource("bridge_terms_2026-09-28.json"))), [])
    assert len(summary) == 1
    assert {x["finding"] for x in findings} >= {"e-mail address", "Old claim ladder (allows 'shows')"}
    assert any("Google Docs shortcut" in x["reason"] for x in skipped)


def test_empty_document_is_not_reported_clean(tmp_path):
    import docx
    from clockbind.docaudit import audit, load_terms
    docx.Document().save(tmp_path / "empty.docx")
    fs, sm, sk = audit(str(tmp_path), load_terms(None), [])
    assert not sm and sk and "no extractable text" in sk[0]["reason"]
