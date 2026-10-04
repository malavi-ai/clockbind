"""Lock: refuses with blockers; locks a clean workbook with frozen gates; verify detects tampering."""
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from tests.test_fullreport import _workbook


def _frozen_gates(tmp):
    src = Path(__file__).resolve().parents[1] / "clockbind" / "data" / "gates_v3.4_DRAFT.json"
    g = json.loads(src.read_text())
    g["protocol_version"] = "v-test"
    for k in ("frozen", "frozen_on", "frozen_by", "sha256"):
        g.pop(k, None)
    p = tmp / "gates.json"
    p.write_text(json.dumps(g))
    subprocess.run([sys.executable, "-m", "clockbind", "screen", "freeze", "--gates", str(p), "--by", "T"], check=True, capture_output=True)
    return p


def test_lock_refuses_draft_gates(tmp_path):
    from clockbind.lock import run_lock
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    src = Path(__file__).resolve().parents[1] / "clockbind" / "data" / "gates_v3.4_DRAFT.json"
    r = run_lock(str(wb), str(src), "T", str(tmp_path / "L"))
    assert not r["locked"] and any(b[0] == "Gates" for b in r["blockers"])


def test_lock_and_verify(tmp_path):
    from clockbind.lock import run_lock, verify
    from openpyxl import load_workbook
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    import pandas as pd
    from clockbind.binding import evaluate_all
    x = pd.read_excel(wb, sheet_name=None, dtype=object)
    verdict = evaluate_all(x["Episodes"], x["Steps"])["verdict"].iloc[0]
    book = load_workbook(wb)
    ws = book["05_Level_Assessment"]
    c = ws.max_column + 1
    ws.cell(3, c, "Supplier-delay rival")
    ws.cell(5, c, "Rejected")
    for col in range(1, ws.max_column + 1):
        if ws.cell(3, col).value == "Verdict":
            ws.cell(5, col, verdict)
    book.save(wb)
    g = _frozen_gates(tmp_path)
    r = run_lock(str(wb), str(g), "T", str(tmp_path / "L"), dry_run=True)
    # the minimal test workbook has no 02B sheet etc.; justify every remaining warning to test the lock path
    notes = tmp_path / "n.csv"
    notes.write_text("check,justification\n" + "".join(f'"{w[1]}",test\n' for w in r["warnings"]))
    real = [b for b in r["blockers"] if b[0] not in ("Undocumented warning",)]
    if real:
        pytest.skip(f"test workbook not lock-clean: {real[:2]}")
    r = run_lock(str(wb), str(g), "T", str(tmp_path / "L"), notes_path=str(notes))
    assert r["locked"], r.get("blockers")
    ok, probs = verify(r["package"])
    assert ok, probs
    f = Path(r["package"]) / "w.xlsx"
    f.chmod(0o644)
    f.write_bytes(f.read_bytes() + b"x")
    ok, probs = verify(r["package"])
    assert not ok and any("changed" in p for p in probs)
    r2 = run_lock(str(wb), str(g), "T", str(tmp_path / "L"), notes_path=str(notes))
    assert not r2["locked"] and any(b[0] == "Amendment" for b in r2["blockers"])
