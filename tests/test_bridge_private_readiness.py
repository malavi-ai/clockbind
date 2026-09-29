from pathlib import Path
import json

from clockbind.docaudit import audit, load_terms, personal_kinds


def test_name_label_does_not_treat_customer_side_hyphen_as_name():
    text = "Level 3 customer-side dates need an outside source; supplier delay is a rival explanation."
    assert personal_kinds(text, [], []) == {}


def test_approved_materiality_shorthand_is_not_flagged(tmp_path):
    terms = load_terms(str(Path(__file__).resolve().parents[1] / "clockbind" / "data" / "bridge_terms_2026-09-28.json"))
    p = tmp_path / "x.txt"
    p.write_text("HUF 200k-equivalent threshold is study-specific.", encoding="utf-8")
    finds, _, _ = audit(str(p), terms, [])
    assert not [x for x in finds if x["check"] == "personal or confidential data" and x["finding"] == "exact amount"]


def test_archive_excel_sheets_are_wording_history_not_current_findings(tmp_path):
    import openpyxl
    terms = load_terms(str(Path(__file__).resolve().parents[1] / "clockbind" / "data" / "bridge_terms_2026-09-28.json"))
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Archive_History"
    ws["A1"] = "four core episodes"
    live = wb.create_sheet("Current")
    live["A1"] = "four core episodes"
    wb.save(p)
    finds, _, _ = audit(str(p), terms, [])
    wording = [x for x in finds if x["check"] == "outdated wording"]
    assert len(wording) == 1
    assert wording[0]["location"].startswith("Current!")
