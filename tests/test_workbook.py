"""Workbook integrity checker on small synthetic workbooks."""
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from clockbind.workbook_check import check_workbook


def _wb(tmp_path, bad_entry=False, override_without_reason=False):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "03_Source_Extraction"
    s.append(["title"]); s.append([]); s.append(["Source code", "x", "y", "z", "Distinct episode extraction status", "Number of distinct episodes", "Stable Episode ID(s)"])
    s.append(["SRC-001", "", "", "", "COMPLETE", 1, "E01"])
    s.append(["SRC-002", "", "", "", "Complete" if bad_entry else "PENDING", None, ""])
    dv = DataValidation(type="list", formula1='"PENDING,COMPLETE,Held"', showErrorMessage=True); dv.add("E4:E5"); s.add_data_validation(dv)
    a = wb.create_sheet("05_Level_Assessment")
    a.append(["t"]); a.append([]); a.append(["Episode ID", "Source code", "Episode extraction complete?", "Override level (only with logged reason)", "Override reason (required if override)", "Override date"])
    a.append(["E01", "SRC-001", "COMPLETE", "Level 1" if override_without_reason else None, None, None])
    p = tmp_path / "w.xlsx"
    wb.save(p)
    return p


def _levels(rep, check):
    return [r["level"] for r in rep.rows if r["check"].startswith(check)]


def test_clean_workbook_passes(tmp_path):
    rep = check_workbook(_wb(tmp_path))
    assert rep.count("FAIL") == 0
    assert "PASS" in _levels(rep, "Every episode code")


def test_dropdown_violation_detected(tmp_path):
    rep = check_workbook(_wb(tmp_path, bad_entry=True))
    assert "FAIL" in _levels(rep, "Entries outside the dropdown list")


def test_override_needs_reason_and_date(tmp_path):
    rep = check_workbook(_wb(tmp_path, override_without_reason=True))
    assert "FAIL" in _levels(rep, "Level overrides have reason and date")


def test_missing_assessment_row_detected(tmp_path):
    p = _wb(tmp_path)
    wb = openpyxl.load_workbook(p)
    wb["03_Source_Extraction"]["G4"] = "E01; E20"
    wb.save(p)
    rep = check_workbook(p)
    fails = [r for r in rep.rows if r["check"].startswith("Every episode code") and r["level"] == "FAIL"]
    assert fails and "E20" in fails[0]["where"] + fails[0]["detail"]
