"""Studio shell: profiles, three languages and the project status."""
import json
from pathlib import Path

import pytest

from clockbind import profiles
from clockbind.studio.i18n import FIX, LANGS, T, digits, fix_text, t

ROOT = Path(__file__).resolve().parents[1]


def test_every_ui_string_in_three_languages():
    for key, row in T.items():
        for L in LANGS:
            assert row.get(L), f"{key} missing {L}"
    for key, row in FIX.items():
        for L in LANGS:
            assert row.get(L) and len(row[L]) == 2, f"fix {key} missing {L}"


def test_placeholders_match_across_languages():
    import re
    for key, row in T.items():
        ph = {L: set(re.findall(r"\{(\w+)\}", row[L])) for L in LANGS}
        assert ph["en"] == ph["hu"] == ph["fa"], key
        if ph["en"]:
            t(key, "fa", **{k: "x" for k in ph["en"]})


def test_persian_digits_and_fallback():
    assert digits("45/19", "fa") == "۴۵/۱۹"
    assert digits("45", "en") == "45"
    assert fix_text("no-such-check", "hu")[0]


def test_profiles_local_roundtrip(tmp_path, monkeypatch):
    monkeypatch.setenv("CLOCKBIND_HOME", str(tmp_path))
    assert profiles.load() == {"profiles": [], "last": None}
    profiles.upsert({"name": "Mr Alavi", "lang": "fa", "workbook": "", "gates": ""})
    profiles.upsert({"name": "Guest", "lang": "hu", "workbook": "", "gates": ""})
    profiles.upsert({"name": "Mr Alavi", "lang": "en", "workbook": "x.xlsx", "gates": ""})
    d = json.loads((tmp_path / "profiles.json").read_text(encoding="utf-8"))
    assert [p["name"] for p in d["profiles"]] == ["Guest", "Mr Alavi"] and d["last"] == "Mr Alavi"
    assert profiles.get("Mr Alavi")["lang"] == "en"
    assert "password" not in (tmp_path / "profiles.json").read_text()
    (tmp_path / "profiles.json").write_text("{broken", encoding="utf-8")
    assert profiles.load()["profiles"] == []


def test_project_status_on_synthetic_workbook():
    from clockbind.project import project_status
    wb = ROOT / "examples" / "screening"
    cands = sorted(wb.glob("*.xlsx"))
    if not cands:
        pytest.skip("no synthetic screening workbook")
    st = project_status(str(cands[0]), str(ROOT / "examples/screening/gates_v3.4_DRAFT.json"))
    assert st["step"] in "ABCDE" and "sources" in st and "levels" in st


def test_studio_opens_and_navigates(tmp_path, monkeypatch):
    AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
    monkeypatch.setenv("CLOCKBIND_HOME", str(tmp_path))
    for L in LANGS:
        at = AppTest.from_file(str(ROOT / "clockbind/studio/app.py"), default_timeout=120).run()
        assert not at.exception
        at.text_input(key="np_name").input("Tester " + L)
        at.selectbox(key="np_lang").select(L)
        at.button[-1].click().run()
        assert not at.exception
        for k in ("nav_bridge", "nav_stats", "nav_document_audit", "nav_privacy", "nav_reproducibility", "nav_settings", "nav_home"):
            at.button(key=k).click().run()
            assert not at.exception, (L, k)


def test_document_audit_finds_data_and_wording(tmp_path):
    import docx
    from clockbind.docaudit import audit, load_names, load_terms
    from clockbind.resources import resource
    d = docx.Document()
    d.add_paragraph("Four episodes were retained in the frozen focal architecture. Grade A permits “shows”.")
    d.add_paragraph("Write to jane@example.com, IBAN GB82 WEST 1234 5698 7654 32, invoice no INV-2026-0042, EUR 12,500. Jane Example agreed.")
    d.add_paragraph("Materiality floor HUF 200,000. ISBN 9781506336169. https://doi.org/10.1287/opre.9.3.296. After the protocol is frozen, coding starts.")
    f = tmp_path / "a.docx"; d.save(f)
    (tmp_path / "names.txt").write_text("Jane Example\n", encoding="utf-8")
    fs, sm, sk = audit(str(tmp_path), load_terms(str(resource("bridge_terms_2026-09-28.json"))), load_names(str(tmp_path / "names.txt")))
    kinds = {x["finding"] for x in fs if x["check"] != "outdated wording"}
    assert kinds == {"e-mail address", "bank account (IBAN)", "invoice or order number", "exact amount", "listed name"}
    assert all(x["matched"] == "" for x in fs if x["check"] != "outdated wording")       # values never reported
    wording = {x["matched"] for x in fs if x["check"] == "outdated wording"}
    assert "four episodes" in {w.lower() for w in wording} and any("shows" in w for w in wording)
    assert not any("frozen" == w.split()[-1] and "protocol" in w for w in wording)       # 'after the protocol is frozen' is fine
    assert "jane" not in str(fs).lower().replace("jane@", "")                            # the listed name never appears


def test_document_audit_reads_pdf_and_zip(tmp_path):
    import zipfile
    from reportlab.pdfgen import canvas
    from clockbind.docaudit import audit, load_terms
    pdf = tmp_path / "b.pdf"
    c = canvas.Canvas(str(pdf)); c.drawString(72, 720, "Contact: someone@example.org. The claim ladder says A = shows."); c.save()
    z = tmp_path / "drive.zip"
    with zipfile.ZipFile(z, "w") as zz:
        zz.write(pdf, "Bridge/b.pdf"); zz.writestr("Bridge/x.gdoc", "{}")
    from clockbind.resources import resource
    fs, sm, sk = audit(str(z), load_terms(str(resource("bridge_terms_2026-09-28.json"))), [])
    assert {x["finding"] for x in fs} >= {"e-mail address", "Old claim ladder (allows 'shows')"}
    assert sk and "Google Docs" in sk[0]["reason"]


def test_research_studio_information_architecture_present():
    shell = (ROOT / "clockbind/studio/shell.py").read_text(encoding="utf-8")
    app = (ROOT / "clockbind/studio/app.py").read_text(encoding="utf-8")
    for key in ("document_audit", "privacy", "reproducibility", "stats", "bridge"):
        assert f'"{key}"' in shell
    for label in ("Upload", "Configure", "Run", "Review", "Export"):
        assert label in shell or label in app
    assert "Create AI-safe research export" in shell
