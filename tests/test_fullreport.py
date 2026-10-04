"""Full report: every output is written, and chat-safe mode carries no episode codes."""
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")


def _workbook(path: Path):
    import json
    from openpyxl import Workbook
    g = json.loads((Path(__file__).resolve().parents[1] / "clockbind" / "data" / "gates_v3.4_DRAFT.json").read_text())
    cols = ["Episode ID", "Source code"] + [c["column"] for s in g["stages"] for c in s["criteria"]] + ["Verdict"]
    wb = Workbook()
    ws = wb.active
    ws.title = "05_Level_Assessment"
    ws.append(["title"]); ws.append(["sub"]); ws.append(cols)
    for ep, fail in (("ZZ1", "L2c"), ("ZZ2", None)):
        row = [ep, "S1"]
        for s in g["stages"]:
            for c in s["criteria"]:
                v = "COMPLETE" if c["id"] == "S0p1" else ("No" if c["id"] == fail else "Yes")
                if c["id"] == "L1f":
                    v = "Small"
                if c["id"] == "L1g":
                    v = "Device/system"
                row.append(v)
        row.append("")
        ws.append(row)
    e = wb.create_sheet("Episodes")
    e.append(["episode", "anchor_earliest", "anchor_latest", "tw_earliest", "tw_latest", "r_step", "note"])
    e.append(["ZZ2", "2025-01-01", "", "2025-01-10", "", "open", ""])
    s = wb.create_sheet("Steps")
    s.append(["episode", "step", "clock", "predecessors", "finish_earliest", "finish_latest", "expected_days", "source", "note"])
    s.append(["ZZ2", "funds", "finance", "", "2025-01-05", "", 2, "D1", ""])
    s.append(["ZZ2", "open", "operational-readiness", "funds", "2025-01-12", "", 2, "D1", ""])
    wb.save(path)


@pytest.mark.parametrize("safe", [False, True])
def test_full_report(tmp_path, safe):
    from clockbind.fullreport import build
    from clockbind.plugins.report import DEFAULT_GATES
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    r = build(str(wb), str(DEFAULT_GATES), str(tmp_path / "out"), chat_safe=safe)
    for k in ("pdf", "xlsx", "docx", "md", "bundle"):
        assert Path(r[k]).exists() and Path(r[k]).stat().st_size > 0
    assert list((tmp_path / "out" / "figures").glob("*.png"))
    text = "".join(p.read_text(errors="ignore") for p in (tmp_path / "out" / "tables").glob("*.csv")) + r["summary"]
    assert ("ZZ1" in text) is (not safe)


def test_consistency_flags(tmp_path):
    from openpyxl import load_workbook
    from clockbind.fullreport import build
    from clockbind.plugins.report import DEFAULT_GATES
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    x = load_workbook(wb)
    ws = x["05_Level_Assessment"]
    ws.cell(5, ws.max_column, "Indeterminate")          # ZZ2 recorded verdict differs from engine
    x["Steps"]["G2"].value = None                        # blank ex-ante duration while L3b = Yes
    x.save(wb)
    r = build(str(wb), str(DEFAULT_GATES), str(tmp_path / "out"))
    c = r["tables"]["11_Consistency_checks"].set_index("check")
    assert c.loc["L3b = Yes in 05, but expected_days blank in Steps", "level"] == "FAIL"
    assert c.loc["Verdict recorded in 05 differs from the binding engine", "level"] == "FAIL"
    assert r["citable"] is False


def test_sensitivity_s1_s2(tmp_path):
    """S1 accepts manual and system records for L3e, S2 only system records; primary levels unchanged."""
    import json as _j
    from openpyxl import load_workbook
    from clockbind.fullreport import build, SENS_COL
    from clockbind.plugins.report import DEFAULT_GATES
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    g = _j.loads(Path(DEFAULT_GATES).read_text())
    l3e = next(c["column"] for s in g["stages"] for c in s["criteria"] if c["id"] == "L3e")
    book = load_workbook(wb)
    ws = book["05_Level_Assessment"]
    hdr = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1)}
    sc = ws.max_column + 1
    ws.cell(3, sc, SENS_COL)
    ws.cell(5, hdr[l3e], "No")                       # ZZ2 now fails L3e only
    ws.cell(5, sc, "standard transactional (manual)")
    book.save(wb)
    t = build(str(wb), str(DEFAULT_GATES), str(tmp_path / "o"))["tables"]["12_Sensitivity_L3e"].set_index("analysis")
    assert t.loc["Primary", "Level 3"] == 0
    assert t.loc["S1", "Level 3"] == 1
    assert t.loc["S2", "Level 3"] == 0


def test_descriptive_fields_ladders_and_safe_share(tmp_path):
    from openpyxl import load_workbook
    from clockbind.fullreport import build
    from clockbind.plugins.report import DEFAULT_GATES
    import json as _j
    wb = tmp_path / "w.xlsx"
    _workbook(wb)
    g = _j.loads(Path(DEFAULT_GATES).read_text())
    col = {c["id"]: c["column"] for s in g["stages"] for c in s["criteria"]}
    book = load_workbook(wb)
    ws = book["05_Level_Assessment"]
    hdr = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1)}
    n = ws.max_column
    for i, name in enumerate(["L3c gap reason", "Buyer stage"], 1):
        ws.cell(3, n + i, name)
    ws.cell(5, hdr[col["L3c"]], "No")                     # ZZ2 fails L3c
    ws.cell(5, n + 1, "informal channel (whatsapp/phone)")  # case-insensitive match
    ws.cell(4, n + 2, "Existing business")
    ws.cell(5, n + 2, "Spaceship")                         # invalid value
    book.save(wb)
    r = build(str(wb), str(DEFAULT_GATES), str(tmp_path / "o"))
    t = r["tables"]
    assert t["13_L3c_gap_reason"].set_index("value").loc["Informal channel (WhatsApp/phone)", "episodes"] == 1
    c = t["11_Consistency_checks"].set_index("check")
    assert c.loc["Descriptive field 'Buyer stage': value not in the allowed list", "level"] == "FAIL"
    lad = t["05_Temporal_ladder"].set_index("step")
    assert lad.filter(like="L3c", axis=0)["No"].iloc[0] == 1
    assert "07_Finance_ladder" in t
    # chat-safe: no codes, manifest written, small cells suppressed
    r2 = build(str(wb), str(DEFAULT_GATES), str(tmp_path / "s"), chat_safe=True)
    assert (tmp_path / "s" / "SAFE_TO_SHARE.txt").exists()
    assert "<3" in r2["tables"]["13_L3c_gap_reason"]["episodes"].astype(str).tolist()
    st = r2["tables"]["00_Status"].set_index("item")
    assert st.loc["Share status", "value"].startswith("SAFE TO SHARE")


def test_leak_detector():
    import pandas as pd
    from clockbind.descriptive import leaked_codes
    assert leaked_codes({"t": pd.DataFrame({"a": ["fine", "X02 mentioned"]})}, ["X02", "X07"]) == ["X02"]
    assert leaked_codes({"t": pd.DataFrame({"a": ["X020"]})}, ["X02"]) == []


def test_claim_tiers_follow_frozen_rule():
    from clockbind.fullreport import claim_tier, claim_wording
    # positive binding: A/B/C by strength of evidence
    assert claim_tier("Level 3", "Finance-binding", "Yes", "Rejected", "A", "Buyer", "Yes") == "A"
    assert claim_tier("Level 3", "Finance-binding", "Yes", "Not rejected", "A", "Buyer", "Yes") == "C"
    # non-binding: negated form at the same tier, never a positive tier
    t = claim_tier("Level 3", "Non-binding", "Yes", "Rejected", "B", "Buyer", "Yes")
    assert t == "B¬" and claim_wording(t) == "appears not to have been binding"
    # indeterminate or sign-unstable -> X; no verdict -> PENDING
    assert claim_tier("Level 3", "Indeterminate", "Yes", "Rejected", "A", "Buyer", "Yes") == "X"
    assert claim_tier("Level 3", "Finance-binding", "No", "Rejected", "A", "Buyer", "Yes") == "X"
    assert claim_tier("Level 3", "", "", "", "", "", "") == "PENDING"
    assert claim_tier("Level 2", "", "", "", "", "", "") == "D"
